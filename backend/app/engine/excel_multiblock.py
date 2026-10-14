"""Plan and run a multi-block Excel report.

`excel_reports.plan_report` binds ONE sheet to ONE query. Real templates are not
shaped that way: CreditFund_RiskDashboard needs three queries across two sheets,
and its other two sheets must receive no data at all because the workbook
computes them itself.

This module is the orchestration over excel_blocks (what is where) and
excel_reflow (make it fit). SQL execution is INJECTED — the caller passes a
function — so the whole pipeline is testable without Snowflake, and so the real
path can keep going through the existing validate_sql gate unchanged.

A note on trust. Block detection is structural but heuristic, and a template it
reads wrongly would produce a confidently wrong workbook — the worst failure a
report can have. So `plan` reports what it found AND what it is unsure about,
and `run` refuses on unresolved problems rather than guessing.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Callable

import openpyxl

from app.engine import excel_blocks, excel_reflow, excel_verify
from app.engine import sql_params

# A query returning far more rows than the template was built around is usually
# a missing WHERE clause rather than genuine growth. Generating it anyway means
# a 40,000-row workbook nobody asked for.
#
# Only meaningful for an INJECTED executor that applies no cap of its own — a
# test, a script. The real executor caps at EXCEL_MAX_ROWS, so this must track
# that number rather than be set independently: when the two were both a
# hard-coded 10,000 the guard could never fire, and when EXCEL_MAX_ROWS was
# raised on its own this one started refusing legitimate reports with a message
# about unfiltered queries. Neither is a judgement anyone made on purpose.
DEFAULT_MAX_GROWTH_ROWS = None


def _growth_cap(explicit: int | None) -> int:
    if explicit is not None:
        return explicit
    from app.engine import excel_reports
    return excel_reports.EXCEL_MAX_ROWS


class MultiBlockError(Exception):
    """A template or binding problem worth reporting to the caller verbatim."""


def snowflake_executor(row_cap: int | None = None) -> Callable[[str], dict[str, Any]]:
    """The real executor: every block's SQL through the SAME read-only gate as
    any other query in the product (validate_sql — SELECT/WITH/UNION only),
    capped at the Excel row limit rather than the 100-row chat limit.

    Kept behind a factory so `run` stays injectable and the pipeline remains
    testable without a Snowflake connection.

    Reports `row_cap` and, when the cap was reached, the true `total_rows` from
    a pre-cap COUNT(*). Without that the caller cannot tell a query that
    genuinely returned exactly `cap` rows from one that was cut off — and a
    report showing the first 10,000 of two million rows, with the workbook's
    totals computed over them, is wrong in a way nobody looking at it can see.
    """
    from app.engine import excel_reports
    from app.engine.sql_safety import enforce_row_limit, validate_sql
    from app.engine.connectors.snowflake import client as snowflake_client

    cap = row_cap or excel_reports.EXCEL_MAX_ROWS

    def execute(sql: str) -> dict[str, Any]:
        safe = enforce_row_limit(validate_sql(sql), cap)
        result = snowflake_client._execute_select(safe)
        result["row_cap"] = cap
        if len(result["rows"]) >= cap:
            # Only when it matters: one extra COUNT(*), for a query we already
            # know may have been truncated.
            result["total_rows"] = excel_reports._count_total_rows(sql)
        return result

    return execute


def _open(source: str | Path | bytes, read_only: bool = False):
    """Load a workbook from a path OR raw bytes.

    Bytes are the important half: once templates live in S3 there is no path to
    hand openpyxl, and a function that only takes paths would have to be
    rewritten rather than reconfigured.
    """
    return openpyxl.load_workbook(
        io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source,
        data_only=False,
        read_only=read_only,
    )


def plan(source: str | Path | bytes,
         role_overrides: dict[str, str] | None = None,
         block_overrides: dict[str, bool] | None = None) -> dict[str, Any]:
    """Read a workbook and report what a report over it would need.

    Returns the blocks a query must be supplied for, the sheets that need none,
    and any structure this module could not read confidently.
    """
    # NOT read_only, despite load_workbook being 15x faster that way (71s -> 5s
    # on a 67k-row feed). A read-only worksheet re-reads the sheet on every
    # ws.cell() call, and detect_blocks makes millions of them — planning went
    # from 75 seconds to over ten minutes. Read-only only pays once detection is
    # rewritten as a single iter_rows pass; the flag alone is a trap.
    wb = _open(source)
    try:
        scan = excel_blocks.scan_blocks(wb, role_overrides, block_overrides)
        features = excel_verify.inventory(wb, source)
        warnings = _structural_warnings(scan) + features["warnings"]
        return {
            "sheets": [
                {
                    "sheet": s["sheet"],
                    "role": s["role"],
                    "reason": s["reason"],
                    "overridden": s.get("overridden", False),
                    "blocks": len(s["blocks"]),
                }
                for s in scan["sheets"]
            ],
            "blocks_to_bind": [
                {
                    "key": f"{b['sheet']}!{b['first_row']}",
                    "sheet": b["sheet"],
                    "title": b["title"],
                    "first_row": b["first_row"],
                    "last_row": b["last_row"],
                    "rows": b["rows"],
                    "spare_rows": b["spare_rows"],
                    "columns_needed": [
                        c["header"] for c in b["columns"]
                        if c["source"] == "value" and c["header"]
                    ],
                    "columns_computed_by_template": [
                        c["column"] for c in b["columns"] if c["source"] == "formula"
                    ],
                }
                for b in scan["bindable_blocks"]
            ],
            "unbound_areas": [
                {
                    "key": f"{b['sheet']}!{b['first_row']}",
                    "sheet": b["sheet"],
                    "title": b["title"],
                    "first_row": b["first_row"],
                    "last_row": b["last_row"],
                    "rows": b["rows"],
                    "why": (
                        "the workbook computes this sheet"
                        if b["sheet_role"] != "source"
                        else "mostly formula columns — looks calculated, not stored"
                    ),
                }
                for b in scan["unbound_blocks"]
            ],
            "summary": scan["summary"],
            "features": features["features"],
            "warnings": warnings,
        }
    finally:
        wb.close()


def _structural_warnings(scan: dict[str, Any]) -> list[str]:
    """Things a human should look at before trusting the binding."""
    out: list[str] = []

    if not scan["bindable_blocks"]:
        out.append(
            "No block on any sheet looks like a query target. Either every sheet "
            "is computed from another, or the tables are too narrow to detect "
            "(a block needs at least three named columns)."
        )

    for s in scan["sheets"]:
        for b in s["blocks"]:
            unnamed = [c for c in b["columns"]
                       if c["source"] == "value" and not c["header"]]
            if unnamed and b.get("needs_query"):
                out.append(
                    f"{b['sheet']}!{b['first_row']}: {len(unnamed)} value column(s) "
                    f"have no header, so nothing can be matched to them by name."
                )
            if b.get("needs_query") and b["rows"] < 2:
                out.append(
                    f"{b['sheet']}!{b['first_row']}: only {b['rows']} sample row(s) — "
                    f"too few to tell a data block from a one-off label."
                )
    return out


def run(
    source: str | Path | bytes,
    block_sql: dict[str, str],
    execute_sql: Callable[[str], dict[str, Any]],
    parameters: list[dict[str, Any]] | None = None,
    parameter_values: dict[str, Any] | None = None,
    max_growth_rows: int | None = DEFAULT_MAX_GROWTH_ROWS,
    allow_unmatched_columns: bool = False,
    allow_truncated_rows: bool = False,
    role_overrides: dict[str, str] | None = None,
    block_overrides: dict[str, bool] | None = None,
) -> dict[str, Any]:
    """Generate the workbook.

    `block_sql` maps a block key ("LOAN_DATA!3") to its parameterised SQL.
    `execute_sql` runs one statement and returns {columns, rows}.

    Returns the workbook as BYTES rather than writing it. Where a report is
    kept — local disk today, S3 next — is the caller's business, and an engine
    that saves to a path can only ever run on a filesystem.
    """
    growth_cap = _growth_cap(max_growth_rows)
    wb = _open(source)
    try:
        scan = excel_blocks.scan_blocks(wb, role_overrides, block_overrides)
        all_blocks = [b for s in scan["sheets"] for b in s["blocks"]]
        bindable = {f"{b['sheet']}!{b['first_row']}": b for b in scan["bindable_blocks"]}

        unknown = set(block_sql) - set(bindable)
        if unknown:
            raise MultiBlockError(
                f"No such block: {', '.join(sorted(unknown))}. "
                f"Known blocks: {', '.join(sorted(bindable))}."
            )
        unbound = set(bindable) - set(block_sql)
        if unbound:
            raise MultiBlockError(
                f"No query supplied for {', '.join(sorted(unbound))}. Every block "
                f"that takes data needs one, or its rows would be left as the "
                f"template's sample values and read as real."
            )

        # --- execute, and learn how many rows each block will hold -----------
        results: dict[str, dict[str, Any]] = {}
        truncated: list[str] = []
        for key, sql_template in block_sql.items():
            rendered = sql_params.render_sql(
                sql_template, parameters or [], parameter_values
            )
            result = execute_sql(rendered)
            n = len(result["rows"])
            if n > growth_cap:
                raise MultiBlockError(
                    f"{key} returned {n:,} rows (cap {growth_cap:,}). That is "
                    f"almost certainly an unfiltered query rather than real growth."
                )
            # Did the executor's own row cap cut this result off? This check used
            # to be impossible: the executor capped at EXCEL_MAX_ROWS and the
            # guard above compared against the SAME number, so it could never
            # fire, and a truncated block sailed through silently.
            cap = result.get("row_cap")
            if cap and n >= cap:
                total = result.get("total_rows")
                truncated.append(
                    f"{key}: {n:,} rows written"
                    + (f" of {total:,} that match" if total else " — the row cap was reached")
                )
            results[key] = {**result, "sql": rendered}

        if truncated and not allow_truncated_rows:
            raise MultiBlockError(
                "The row cap cut these queries short:\n  "
                + "\n  ".join(truncated)
                + "\nThe workbook's own totals would be computed over the rows "
                  "that fit and would read as if they covered everything. Narrow "
                  "the query, raise MCP_EXCEL_MAX_ROWS, or re-run allowing "
                  "truncation if a partial extract is what you want."
            )

        # --- column matching, before anything is written ---------------------
        column_maps: dict[str, dict[str, str]] = {}
        problems: list[str] = []
        for key, block in bindable.items():
            cmap = excel_reflow.match_columns(block, results[key]["columns"])
            missing = excel_reflow.unmatched_headers(block, cmap)
            column_maps[key] = cmap
            if missing:
                problems.append(
                    f"{key}: the query supplies nothing for {', '.join(missing)}"
                )
        if problems and not allow_unmatched_columns:
            raise MultiBlockError(
                "Some template columns would be left blank:\n  "
                + "\n  ".join(problems)
                + "\nAlias the query's columns to match the template's headers, or "
                  "re-run allowing unmatched columns if the blanks are intended."
            )

        # --- row counts, including blocks that mirror a source ---------------
        requested = {
            (b["sheet"], b["first_row"]): len(results[k]["rows"])
            for k, b in bindable.items()
        }
        new_rows = excel_reflow.link_mirror_blocks(wb, all_blocks, requested)

        # --- reflow every sheet, then fill ----------------------------------
        by_sheet = {s["sheet"]: s for s in scan["sheets"]}
        row_maps = {
            name: excel_reflow.build_row_map(
                s["blocks"],
                {b["first_row"]: new_rows.get((name, b["first_row"]), b["rows"])
                 for b in s["blocks"]},
                wb[name].max_row,
            )
            for name, s in by_sheet.items()
        }

        def row_map_for(sheet: str):
            return row_maps.get(sheet)

        moved = 0
        for name, s in by_sheet.items():
            moved += excel_reflow.reflow_sheet(
                wb[name], s["blocks"], {}, row_maps[name], row_map_for
            )["cells_moved"]

        filled = []
        for key, block in bindable.items():
            res = results[key]
            filled.append(excel_reflow.fill_block(
                wb[block["sheet"]], block, row_maps[block["sheet"]],
                res["columns"], res["rows"], column_maps[key],
            ))

        regenerated = []
        for b in all_blocks:
            k = (b["sheet"], b["first_row"])
            if k in requested or k not in new_rows:
                continue
            r = excel_reflow.regenerate_block(
                wb[b["sheet"]], b, row_maps[b["sheet"]], new_rows[k]
            )
            if r["rows"]:
                regenerated.append(r)

        # The workbook ships formulas with no cached results; without this,
        # every derived cell opens blank.
        wb.calculation.fullCalcOnLoad = True
        buffer = io.BytesIO()
        wb.save(buffer)
    finally:
        wb.close()

    # Compare the finished file against the template it came from. This is the
    # only check that catches damage to features nobody thought about: whatever
    # was there has to still be there, moved but not lost.
    losses: list[str] = []
    try:
        original = _open(source)
        rebuilt = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
        try:
            losses = excel_verify.conservation(
                original, rebuilt, {b["sheet"] for b in filled}
            )
        finally:
            original.close()
            rebuilt.close()
    except Exception as exc:  # noqa: BLE001 - a failed check must not lose the report
        losses = [f"could not verify the output: {exc}"]

    return {
        "workbook": buffer.getvalue(),
        "blocks_filled": filled,
        "blocks_regenerated": regenerated,
        "cells_moved": moved,
        "row_counts": {k: len(v["rows"]) for k, v in results.items()},
        "column_warnings": problems,
        "truncated_blocks": truncated,
        "lost_content": losses,
    }
