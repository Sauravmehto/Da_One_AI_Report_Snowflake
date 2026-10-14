"""Find the REGIONS in a workbook, and say which of them a query should feed.

The existing scanner classifies whole SHEETS ("data" / "formula" / "parameter"),
which real templates defeat: CreditFund_RiskDashboard has three separate tables
stacked down one sheet, and a data tab holding three computed columns is scored
as "formula" and never offered as a data sheet at all.

This module works one level down. A sheet is a list of BLOCKS — a header row
followed by a contiguous run of records — each with a real extent, so the two
questions that matter can be answered:

  * Which blocks does a query feed?  A block whose sheet is computed FROM another
    sheet needs no query: the workbook already derives it, and writing values
    over those formulas would destroy the recalculation the report exists for.
  * What sits below a block?  A block that grows past the next thing on the sheet
    overwrites it. In this template the loan register has exactly TWO spare rows
    before the FX table, so a query returning more than seven loans silently
    destroys it.

Pure: takes an openpyxl worksheet, returns plain dicts. No Snowflake, no I/O.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter

# A run of records must be at least this tall to count as a block.
#
# ONE, not two. A template is meant to ship blank or near-blank on its raw
# sheets — sample rows left in a delivered workbook read as real data — so
# requiring two of them punished exactly the templates we want people to make.
# One row is still enough to type the columns and spot a fill-down formula.
MIN_BLOCK_ROWS = 1

# A header row must name at least this many columns.
#
# Three, not two: a parameter sheet is a column of labels beside a column of
# values, and every one of its rows reads as a two-column "header" followed by
# more of the same. ASSUMPTIONS was detected as two data tables that way. Real
# record tables in these templates are 5-36 columns wide, so the cost of the
# stricter rule is a genuine two-column table, which is rare and can be mapped
# by hand.
MIN_HEADER_COLS = 3

_SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_.]*))!")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _kind(value: Any) -> str | None:
    """None (empty) / 'formula' / 'value'."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return "formula" if _is_formula(value) else "value"


def referenced_sheets(ws) -> set[str]:
    """Sheet names this sheet's formulas read from (excluding itself)."""
    out: set[str] = set()
    for row in ws.iter_rows():
        for cell in row:
            if _is_formula(cell.value):
                for m in _SHEET_REF.finditer(cell.value):
                    name = m.group(1) or m.group(2)
                    if name and name.upper() != ws.title.upper():
                        out.add(name)
    return out


def _row_profile(ws, row: int, max_col: int) -> dict[str, Any]:
    values = formulas = 0
    first_col = last_col = None
    for c in range(1, max_col + 1):
        k = _kind(ws.cell(row=row, column=c).value)
        if k is None:
            continue
        if first_col is None:
            first_col = c
        last_col = c
        if k == "formula":
            formulas += 1
        else:
            values += 1
    return {
        "row": row,
        "values": values,
        "formulas": formulas,
        "populated": values + formulas,
        "first_col": first_col,
        "last_col": last_col,
    }


def _looks_like_header(ws, row: int, max_col: int) -> bool:
    """A header row is text across several columns, none of them formulas."""
    labels = 0
    for c in range(1, max_col + 1):
        v = ws.cell(row=row, column=c).value
        if _is_formula(v):
            return False
        if isinstance(v, str) and v.strip():
            labels += 1
        elif v is not None:
            return False  # a number in the header row — probably data
    return labels >= MIN_HEADER_COLS


def _section_title(ws, above_row: int, max_col: int) -> str | None:
    """The nearest single-cell banner above a header row ('A. Loan-Level …')."""
    for r in range(above_row - 1, max(0, above_row - 4), -1):
        prof = _row_profile(ws, r, max_col)
        if prof["populated"] == 1 and prof["first_col"] == 1:
            text = _text(ws.cell(row=r, column=1).value)
            if text and not _looks_like_header(ws, r, max_col):
                return text
    return None


def _column_spec(ws, header_row: int, first_row: int, last_row: int,
                 max_col: int) -> list[dict[str, Any]]:
    """Per-column: its header, and whether the block holds values or formulas
    there. A formula column is the template's own calculation and must be
    re-created by filling down, never overwritten with query output."""
    cols: list[dict[str, Any]] = []
    for c in range(1, max_col + 1):
        header = _text(ws.cell(row=header_row, column=c).value)
        kinds = {
            _kind(ws.cell(row=r, column=c).value)
            for r in range(first_row, last_row + 1)
        } - {None}
        if not header and not kinds:
            continue
        is_formula = kinds == {"formula"}
        cols.append({
            "column": get_column_letter(c),
            "header": header or None,
            "source": "formula" if is_formula else "value",
            # The first row's formula is the fill-down template for the block.
            "formula_template": (
                ws.cell(row=first_row, column=c).value if is_formula else None
            ),
        })
    return cols


def detect_blocks(ws) -> list[dict[str, Any]]:
    """Every header-plus-records block on the sheet, in row order."""
    max_row, max_col = ws.max_row, ws.max_column
    profiles = [_row_profile(ws, r, max_col) for r in range(1, max_row + 1)]
    by_row = {p["row"]: p for p in profiles}

    blocks: list[dict[str, Any]] = []
    r = 1
    while r <= max_row:
        prof = by_row[r]
        if prof["populated"] == 0 or not _looks_like_header(ws, r, max_col):
            r += 1
            continue

        header_row = r
        header_width = prof["populated"]
        # Records run until a blank row, or a row that is itself a new header,
        # or a row far narrower than the header (a TOTAL line or a new banner).
        first_row = header_row + 1
        last_row = header_row
        rr = first_row
        while rr <= max_row:
            p = by_row[rr]
            if p["populated"] == 0:
                break
            if _looks_like_header(ws, rr, max_col) and p["populated"] >= MIN_HEADER_COLS:
                break
            # A section banner is one or two cells of prose. Test that in
            # ABSOLUTE terms, never as a fraction of the header: a wide feed is
            # legitimately sparse — 82 populated columns out of 229 is ordinary
            # data, and a "fills half the header" rule discarded a 419-row
            # sheet on its very first record.
            if p["populated"] <= 2 and header_width >= 4:
                break
            # A TOTAL line is nearly as wide as a record but not quite — it
            # skips the label columns it is summarising. Compare against the
            # RECORD rows, not the header: measuring against the header let
            # `TOTAL / WTD AVG` (5 of 9 columns) count as a sixth data row, and
            # the block then grew by one every run.
            if last_row > header_row:
                prev = by_row[last_row]["populated"]
                if p["populated"] < prev * 0.7:
                    break
            last_row = rr
            rr += 1

        if last_row - first_row + 1 >= MIN_BLOCK_ROWS:
            cols = _column_spec(ws, header_row, first_row, last_row, max_col)
            value_cols = [c for c in cols if c["source"] == "value"]
            formula_cols = len(cols) - len(value_cols)
            blocks.append({
                # A block the workbook computes for itself. TIMESERIES's
                # roll-up is nine formula columns over the detail block above
                # it — same sheet, so no cross-sheet reference gives it away,
                # but querying it would overwrite the aggregation.
                "needs_query": len(value_cols) > formula_cols,
                "sheet": ws.title,
                "title": _section_title(ws, header_row, max_col),
                "header_row": header_row,
                "first_row": first_row,
                "last_row": last_row,
                "rows": last_row - first_row + 1,
                "first_col": prof["first_col"],
                "last_col": prof["last_col"],
                "columns": cols,
                "value_columns": len(value_cols),
                "formula_columns": len(cols) - len(value_cols),
            })
            r = last_row + 1
        else:
            r = header_row + 1

    _annotate_room(blocks, max_row)
    return blocks


def _annotate_room(blocks: list[dict[str, Any]], sheet_max_row: int) -> None:
    """How far each block can grow before it lands on something.

    `spare_rows` is the gap to whatever comes next; growing beyond it overwrites
    that content unless the run reflows the rows below.
    """
    for i, b in enumerate(blocks):
        if i + 1 < len(blocks):
            nxt = blocks[i + 1]
            # The next block's banner/header, whichever comes first.
            obstacle = nxt["header_row"]
            if nxt.get("title"):
                obstacle -= 1
            b["next_content_row"] = obstacle
            b["spare_rows"] = max(0, obstacle - b["last_row"] - 1)
        else:
            b["next_content_row"] = None
            b["spare_rows"] = max(0, sheet_max_row - b["last_row"])


def classify_sheet(ws, blocks: list[dict[str, Any]]) -> dict[str, Any]:
    """Whether this sheet is fed by queries or computed by the workbook.

    A sheet whose formulas read another sheet is DERIVED: the template already
    knows how to produce it, and it must be left alone apart from extending the
    ranges its formulas cover. Only a sheet that holds records of its own, with
    no such dependency, needs a query bound to it.
    """
    reads = referenced_sheets(ws)
    has_records = any(b["value_columns"] > 0 for b in blocks)

    if reads:
        role = "computed"
        reason = f"formulas read {', '.join(sorted(reads))} — the workbook derives this sheet"
    elif has_records:
        role = "source"
        reason = "holds records of its own and reads no other sheet"
    else:
        role = "parameter"
        reason = "no record block — labels and values only"

    return {
        "sheet": ws.title,
        "role": role,
        "reason": reason,
        "reads_sheets": sorted(reads),
        "blocks": blocks,
    }


def scan_blocks(wb, role_overrides: dict[str, str] | None = None,
                block_overrides: dict[str, bool] | None = None) -> dict[str, Any]:
    """Whole-workbook block map: what each sheet is, and where a query goes.

    `role_overrides` lets a person correct the detection — {"Sheet1": "source"}.
    Detection is structural but heuristic, and without a way to say "no, this
    one does take data" a wrong guess is a dead end.
    """
    overrides = {k.upper(): v for k, v in (role_overrides or {}).items()}
    # Per-AREA corrections, keyed "SHEET!firstrow". A sheet-level override
    # cannot express this: TIMESERIES already takes data, yet its roll-up table
    # was rejected by the separate per-block test for being mostly formulas —
    # so the area a user actually wanted to bind was simply absent.
    block_flags = {k.upper(): v for k, v in (block_overrides or {}).items()}
    sheets = []
    for ws in wb.worksheets:
        info = classify_sheet(ws, detect_blocks(ws))
        chosen = overrides.get(ws.title.upper())
        if chosen and chosen != info["role"]:
            info["role"] = chosen
            info["reason"] = f"set to '{chosen}' by hand (detected '{info['reason']}')"
            info["overridden"] = True
            # Forcing a sheet to "source" has to overrule the per-block
            # needs_query test too. That test rejects a block whose columns are
            # mostly formulas — which is precisely the guess being corrected, so
            # honouring it here left the override relabelling a row and offering
            # nothing to bind.
            if chosen == "source":
                for b in info["blocks"]:
                    b["needs_query"] = True
        for b in info["blocks"]:
            flag = block_flags.get(f"{ws.title}!{b['first_row']}".upper())
            if flag is not None:
                b["needs_query"] = flag
                b["overridden"] = True
        sheets.append(info)

    bindable = [
        {**b, "sheet_role": s["role"]}
        for s in sheets
        if s["role"] == "source"
        for b in s["blocks"]
        if b["needs_query"]
    ]
    # Areas we did NOT offer, so the UI can show them and let someone turn one
    # on. Silently omitting them is what made a real table look invisible.
    not_bound = [
        {**b, "sheet_role": s["role"]}
        for s in sheets
        for b in s["blocks"]
        if not (s["role"] == "source" and b["needs_query"])
    ]
    return {
        "sheets": sheets,
        "bindable_blocks": bindable,
        "unbound_blocks": not_bound,
        "summary": {
            "sheets": len(sheets),
            "source_sheets": sum(1 for s in sheets if s["role"] == "source"),
            "computed_sheets": sum(1 for s in sheets if s["role"] == "computed"),
            "parameter_sheets": sum(1 for s in sheets if s["role"] == "parameter"),
            "blocks_needing_a_query": len(bindable),
        },
    }
