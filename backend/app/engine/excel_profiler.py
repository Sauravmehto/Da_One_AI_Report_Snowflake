"""Excel data profiler — describe an uploaded workbook so its data can be
extracted, mapped to business meaning, and queried.

This is NOT excel_template_scanner.py and does not replace it. That module
answers "can this template's formulas be safely regenerated?" for the report
builder. This one answers a different question: "what DATA is in this workbook,
where does it live, and what is deterministically knowable about it?"

Both read workbooks; neither imports the other. This module is purely additive —
nothing in the report-builder path changes.

DELIBERATE SPLIT — deterministic here, semantic elsewhere
---------------------------------------------------------
Real client workbooks are not tables. Two samples, two entirely different
shapes, and no rule set would resolve either on its own:

  * An income statement with 20 property tabs, each a *form* — line items down
    column A, values in B, section headers interleaved, subtotals mixed in with
    inputs, and dollars and percentages sharing one column.
  * A rent roll with TWO tables on one sheet whose columns do not align, zero
    data rows (it is a blank template), and total rows carrying literal zeros
    rather than formulas.

So this module extracts only what a machine can know for certain — cell types,
number formats, formula-ness, merged ranges, blank-row boundaries, total-row
labels — and packages it as evidence. Deciding what a column MEANS is left to
the caller (an LLM), because meaning is the part that genuinely needs judgement.

That split is the whole design: structure and units are computed, meaning is
inferred, and the number a user finally sees is always computed by a query
engine over extracted rows — never read out of a cell by a model.

WHY BOTH SUBTOTAL SIGNALS
-------------------------
A row that must be excluded from aggregates announces itself in one of two ways,
and neither alone is sufficient. In the income statement every subtotal is a
formula (`=SUM(B12:B18)`) and every input is a literal. In the rent roll the
total row holds a literal `0.0` and is identifiable only by the word "Total" in
its label column. We therefore test for both and flag a row if either fires.

FORMATS
-------
openpyxl reads .xlsx/.xlsm only. Client files are routinely still legacy OLE2
.xls, which xlrd reads. Both are normalised into the same Grid so that
everything downstream is format-blind. .xls carries less information — xlrd
exposes no formula text — so `Cell.formula` is always None there and subtotal
detection on .xls leans entirely on the label signal.
"""

from __future__ import annotations

import bisect
import hashlib
import re
from datetime import date, datetime
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Literal

import openpyxl

# --- tuning constants (documented so they can be adjusted from real results) ---

#: A header row must be at least this fraction text, ignoring empty cells.
HEADER_TEXT_RATIO = 0.75
#: A records region needs a header plus at least this many rows beneath it
#: before we will call the columns type-consistent. Below this we still report
#: the region (an empty template is legitimate) but mark it low confidence.
MIN_RECORD_ROWS = 2
#: A region whose populated cells are at least this fraction in two columns
#: looks like a label/value form rather than a wide table.
ENTITY_TWO_COL_RATIO = 0.8
#: Blank rows required to end one region and start the next.
REGION_GAP_ROWS = 1

#: Labels that mark a row as an aggregate of other rows. Matched at the start of
#: the label, case-insensitively, so "Total operating expenses" fires.
_TOTAL_LABEL = re.compile(
    r"^\s*(grand\s+)?(total|subtotal|sub-total|sum)\b", re.IGNORECASE
)

#: A section heading inside a form: all-caps text on a row with no value beside
#: it, e.g. "OPERATING EXPENSES".
_SECTION_LABEL = re.compile(r"^[A-Z0-9][A-Z0-9 &/()\-.,']*$")

#: The period a file represents, as spreadsheets actually write it — "as of
#: June 30, 2026", "As Of: 06/30/2026", "for the quarter ended March 31 2026".
#: This is what lets twelve monthly files become a time series instead of twelve
#: unrelated uploads, so it is worth reading rather than asking for.
_AS_OF = re.compile(
    r"(?:as[\s_-]*of|as[\s_-]*at|period[\s_-]*end(?:ing|ed)?|quarter[\s_-]*end(?:ing|ed)?|"
    r"month[\s_-]*end(?:ing|ed)?|for[\s_-]*the[\s_-]*(?:month|quarter|year)[\s_-]*end(?:ing|ed)?)"
    r"\s*[:\-—]?\s*"
    r"(?P<date>"
    r"[A-Z][a-z]{2,8}\.?\s+\d{1,2},?\s+\d{4}"          # June 30, 2026
    r"|\d{1,2}[/-]\d{1,2}[/-]\d{2,4}"                   # 06/30/2026
    r"|\d{4}-\d{2}-\d{2}"                               # 2026-06-30
    r")",
    re.IGNORECASE,
)

#: A bare date in a filename: RentRoll_2026-06.xlsx, Rent Roll Jun 2026.xlsx.
#: Only consulted when the workbook itself does not say.
_FILENAME_DATE = re.compile(
    r"(?P<ym>(?:19|20)\d{2})[-_. ]?(?P<m>0[1-9]|1[0-2])(?:[-_. ]?(?P<d>[0-3]\d))?"
    r"|(?P<mon>jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[-_. ]?(?P<y2>(?:19|20)\d{2})",
    re.IGNORECASE,
)

CellKind = Literal["text", "number", "date", "bool", "error", "empty"]
Unit = Literal["currency", "percent", "date", "number", "text", "unknown"]
RegionKind = Literal["records", "entity", "prose", "unknown"]


# ── primitives ────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Cell:
    """One populated cell, normalised across the two workbook formats."""

    row: int  # 1-based, matching what a user sees in Excel
    col: int  # 1-based
    value: Any
    kind: CellKind
    formula: str | None = None
    number_format: str | None = None

    @property
    def is_derived(self) -> bool:
        """True when this cell computes itself from others, so it must never be
        summed alongside them."""
        return self.formula is not None

    @property
    def unit(self) -> Unit:
        return unit_from_format(self.number_format, self.kind)

    @property
    def currency(self) -> str | None:
        """Which currency this cell's format names, if it names one."""
        if self.unit != "currency":
            return None
        return currency_from_format(self.number_format)


@dataclass
class Grid:
    """A sheet reduced to its populated cells. Sparse on purpose: real sheets
    are mostly empty, and every consumer here asks "what is at (r, c)?" rather
    than iterating a dense rectangle."""

    name: str
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)
    n_rows: int = 0
    n_cols: int = 0
    merged: list[str] = field(default_factory=list)
    #: Row number -> that row's cells, left to right. Built on first use.
    _rows: dict[int, list[Cell]] | None = field(
        default=None, repr=False, compare=False
    )

    def at(self, row: int, col: int) -> Cell | None:
        return self.cells.get((row, col))

    def _index(self) -> dict[int, list[Cell]]:
        """Group cells by row, once.

        row_cells is called several times per row — by the band splitter, the
        header finder, the total-row test and the loader — and it used to sort
        the whole cell dict on every call. On a 240-row sheet with 3,600 cells
        that is a few hundred full sorts per sheet, which dominated the cost of
        an entire upload. Grids are built once and only read afterwards, so a
        single pass is enough.
        """
        if self._rows is None:
            index: dict[int, list[Cell]] = {}
            for (row, _), cell in self.cells.items():
                index.setdefault(row, []).append(cell)
            for row_cells in index.values():
                row_cells.sort(key=lambda c: c.col)
            object.__setattr__(self, "_rows", index)
        return self._rows  # type: ignore[return-value]

    def row_cells(self, row: int) -> list[Cell]:
        return self._index().get(row, [])

    def is_blank_row(self, row: int) -> bool:
        return not self._index().get(row)

    def populated_rows(self) -> list[int]:
        return sorted(self._index())


# ── unit inference ────────────────────────────────────────────────────────────

_PERCENT = re.compile(r"%")
_CURRENCY = re.compile(r"[$£€¥]|\[\$")
_DATEISH = re.compile(r"(yy|mm|dd|hh)", re.IGNORECASE)

#: The currency inside Excel's bracket construct: "[$AED-4001] #,##0.00", or
#: "[$€-407]", or an accounting format's "_-[$£-en-GB]* #,##0.00_-". What
#: sits between "[$" and the first "-" (or the closing bracket) is the symbol
#: or the ISO code; everything after the dash is a locale id and says nothing
#: about which currency it is.
_BRACKET_CURRENCY = re.compile(r"\[\$([^\]\-]*)")

#: Symbols that name one currency clearly enough to record a code for.
#:
#: "$" is deliberately ABSENT. It is the dollar of a dozen countries and the
#: locale id after the dash is the only thing separating them, so reading USD
#: out of a bare "$" would be the very error this module exists to prevent.
#: Leaving it unnamed costs nothing: an unnamed currency already renders as
#: dollars downstream, so a "$" format comes out right without being guessed.
#:
#: The three below are inferences too, honestly: "¥" is also the yuan and "£"
#: is also the Egyptian and Lebanese pound. They are recorded anyway because
#: the code is only ever used to look the symbol back up again -- JPY renders
#: "¥" and GBP renders "£", the same glyph the workbook showed. The alternative
#: is not a safer guess, it is falling back to "$", a different currency.
_SYMBOL_CURRENCY = {"€": "EUR", "£": "GBP", "¥": "JPY"}


def currency_from_format(number_format: str | None) -> str | None:
    """The ISO 4217 code a cell's number format NAMES, or None.

    unit_from_format answers "is this money"; this answers "which money". They
    used to be one function, and that union was a lie by omission: a column of
    dirhams formatted "[$AED-4001] #,##0.00" was correctly recognised as
    currency and then rendered with a dollar sign, because the only place in
    the whole pipeline that knew it was dirhams discarded the code one line
    after matching it. The workbook is the single authoritative statement of
    the currency; nothing downstream can recover it once it is dropped here.

    None means "not named here", NOT "no currency": a bare "$#,##0.00" is a
    currency format that names nothing, and the dollar default renders it
    correctly. Only a code that is actually stated is returned.
    """
    fmt = number_format or ""
    match = _BRACKET_CURRENCY.search(fmt)
    if match:
        inner = match.group(1).strip()
        # "[$AED-4001]" -- a three-letter code, the case that matters most.
        if len(inner) == 3 and inner.isascii() and inner.isalpha():
            return inner.upper()
        return _SYMBOL_CURRENCY.get(inner)

    # No bracket: a bare or quoted symbol, e.g. £#,##0.00
    for symbol, code in _SYMBOL_CURRENCY.items():
        if symbol in fmt:
            return code
    return None


def unit_from_format(number_format: str | None, kind: CellKind) -> Unit:
    """What a cell's number format proves about its unit.

    This is the highest-value deterministic signal in the whole module. In the
    income statement, column B holds dollars in one row and a ratio three rows
    down; no column-level heuristic can be right about both, but each cell's own
    format says exactly which it is.
    """
    if kind == "text":
        return "text"
    if kind == "date":
        return "date"
    if kind in ("bool", "error", "empty"):
        return "unknown"

    fmt = number_format or ""
    if _PERCENT.search(fmt):
        return "percent"
    if _CURRENCY.search(fmt):
        return "currency"
    if _DATEISH.search(fmt):
        return "date"
    return "number"


_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def parse_as_of(text: str) -> date | None:
    """The date in "as of June 30, 2026", or None.

    Deliberately narrow: it only fires on a phrase that ANNOUNCES a period.
    A bare date somewhere on the sheet is far more likely to be a lease
    expiry or a print date, and reading one of those as the file's period
    would file a month's rows under the wrong month — silently, and in a way
    only a trend chart would reveal.
    """
    match = _AS_OF.search(text or "")
    if not match:
        return None
    return _to_date(match.group("date"))


def _to_date(raw: str) -> date | None:
    raw = raw.strip().rstrip(",")
    for fmt in ("%B %d, %Y", "%B %d %Y", "%b %d, %Y", "%b %d %Y",
                "%m/%d/%Y", "%m-%d-%Y", "%Y-%m-%d", "%m/%d/%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def as_of_from_filename(filename: str) -> date | None:
    """A period read off the file name, for workbooks that do not state one.

    Weaker evidence than the sheet, so callers should prefer `parse_as_of`
    and fall back here. A day is assumed to be the 1st when the name carries
    only a month, because a month is the real granularity of these files.
    """
    match = _FILENAME_DATE.search(Path(filename).stem)
    if not match:
        return None
    if match.group("ym"):
        year = int(match.group("ym"))
        month = int(match.group("m"))
        day = int(match.group("d") or 1)
    else:
        year = int(match.group("y2"))
        month = _MONTHS[match.group("mon").lower()[:3]]
        day = 1
    try:
        return date(year, month, day)
    except ValueError:
        return None


def normalise_label(value: Any) -> str:
    """Header and label text as a stable string.

    Real headers carry newlines ("Comments\\n(Add notes as needed)") and stray
    whitespace. Collapsing them keeps a label usable as a dictionary key without
    discarding the parenthetical, which often carries the unit or a hint that the
    field is derived ("Lease Term (Auto Calc Months)").
    """
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


# ── readers ───────────────────────────────────────────────────────────────────


def _kind_from_python(value: Any) -> CellKind:
    import datetime as _dt

    if value is None or value == "":
        return "empty"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return "date"
    if isinstance(value, str) and value.startswith("#") and value.endswith("!"):
        return "error"
    return "text"


def _grids_from_xlsx(path: Path) -> list[Grid]:
    """Load .xlsx/.xlsm twice — once for cached values, once for formula text.

    openpyxl gives one or the other per load, never both. The values load is what
    a user sees; the formula load is what tells us a cell is derived. Every
    sample we have carries cached values, so we never evaluate a formula.
    """
    wb_values = openpyxl.load_workbook(path, data_only=True)
    wb_formulas = openpyxl.load_workbook(path, data_only=False)

    grids: list[Grid] = []
    for ws_v in wb_values.worksheets:
        ws_f = wb_formulas[ws_v.title]
        grid = Grid(
            name=ws_v.title,
            n_rows=ws_v.max_row or 0,
            n_cols=ws_v.max_column or 0,
            merged=[str(r) for r in ws_v.merged_cells.ranges],
        )
        for row in ws_v.iter_rows():
            for cell in row:
                value = cell.value
                raw = ws_f.cell(row=cell.row, column=cell.column).value
                formula = None
                if isinstance(raw, str) and raw.startswith("="):
                    formula = raw
                elif type(raw).__name__ in ("ArrayFormula", "DataTableFormula"):
                    # openpyxl wraps these rather than returning a string; the
                    # text is not what matters, only that the cell is derived.
                    formula = "=<array>"

                blank = value is None or (isinstance(value, str) and not value.strip())
                if blank and formula is None:
                    continue
                if blank:
                    # A formula with no cached value — the workbook has never
                    # been calculated, or was written by a tool that does not
                    # evaluate. The field and its derivation are still real, so
                    # dropping the cell would hide a concept entirely; only the
                    # number is unknown, and derived numbers are excluded from
                    # aggregates regardless.
                    value = None

                grid.cells[(cell.row, cell.column)] = Cell(
                    row=cell.row,
                    col=cell.column,
                    value=value,
                    kind=_kind_from_python(value),
                    formula=formula,
                    number_format=cell.number_format,
                )
        grids.append(grid)
    return grids


def _grids_from_xls(path: Path) -> list[Grid]:
    """Load legacy OLE2 .xls via xlrd.

    `formatting_info=True` is what makes number formats — and therefore units —
    available. xlrd exposes no formula text, so every Cell.formula here is None
    and subtotal detection falls back on the label signal alone.
    """
    import xlrd

    try:
        book = xlrd.open_workbook(path, formatting_info=True)
        has_formats = True
    except NotImplementedError:
        # Some .xls variants refuse formatting_info; values still load.
        book = xlrd.open_workbook(path)
        has_formats = False

    kind_by_ctype: dict[int, CellKind] = {
        xlrd.XL_CELL_EMPTY: "empty",
        xlrd.XL_CELL_TEXT: "text",
        xlrd.XL_CELL_NUMBER: "number",
        xlrd.XL_CELL_DATE: "date",
        xlrd.XL_CELL_BOOLEAN: "bool",
        xlrd.XL_CELL_ERROR: "error",
        xlrd.XL_CELL_BLANK: "empty",
    }

    grids: list[Grid] = []
    for sheet in book.sheets():
        grid = Grid(
            name=sheet.name,
            n_rows=sheet.nrows,
            n_cols=sheet.ncols,
            merged=[],
        )
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                raw = sheet.cell(r, c)
                kind = kind_by_ctype.get(raw.ctype, "text")
                if kind == "empty":
                    continue
                value = raw.value
                if isinstance(value, str):
                    value = value.strip()
                    if not value:
                        continue

                number_format = None
                if has_formats:
                    try:
                        xf = book.xf_list[raw.xf_index]
                        number_format = book.format_map[xf.format_key].format_str
                    except (AttributeError, IndexError, KeyError, TypeError):
                        number_format = None

                grid.cells[(r + 1, c + 1)] = Cell(
                    row=r + 1,
                    col=c + 1,
                    value=value,
                    kind=kind,
                    formula=None,  # xlrd cannot supply it
                    number_format=number_format,
                )
        grids.append(grid)
    return grids


def read_grids(path: str | Path) -> list[Grid]:
    """Every sheet of a workbook, normalised, whichever format it is in."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return _grids_from_xlsx(path)
    if suffix == ".xls":
        return _grids_from_xls(path)
    raise ValueError(
        f"Unsupported workbook format {suffix!r}. "
        "Expected .xlsx, .xlsm, .xltx, .xltm, or .xls."
    )


# ── region detection ──────────────────────────────────────────────────────────


#: A string that is really a number: '$113,990', '100%', '(1,234)', '9.63'.
#: Rent rolls store these as text constantly, and the existing Snowflake
#: semantic layers already carry hand-written notes about parsing exactly this.
_NUMERIC_TEXT = re.compile(
    r"^\s*\(?\s*[-+]?\s*[$£€¥]?\s*\d{1,3}(,\d{3})*(\.\d+)?\s*%?\s*\)?\s*$"
)
#: A string that is really a date: '12/30/2023', '2023-12-30'.
_DATE_TEXT = re.compile(r"^\s*\d{1,4}[-/]\d{1,2}[-/]\d{1,4}\s*$")


@dataclass
class ColumnProfile:
    """One column of a records region."""

    col: int
    header: str
    kinds: dict[CellKind, int] = field(default_factory=dict)
    units: dict[Unit, int] = field(default_factory=dict)
    #: ISO codes seen in this column's cell formats, counted. A dict rather
    #: than a single value because a column CAN mix currencies, and a mixed
    #: column has to stay distinguishable from a single-currency one --
    #: stamping one code over both is how a total in dirhams ends up
    #: labelled euros.
    currencies: dict[str, int] = field(default_factory=dict)
    derived_count: int = 0
    populated: int = 0
    #: Text cells whose contents are really a number or a date. Counted so the
    #: caller knows a cast is required before the column can be aggregated.
    numeric_as_text: int = 0
    date_as_text: int = 0
    #: Cells holding a formula whose result was never stored. Excel writes the
    #: last computed value into the file; a workbook that has not been opened
    #: and calculated — or was opened behind the "Enable content" bar — has the
    #: formulas but none of their results, so the column reads as entirely
    #: empty while looking perfectly well-formed.
    uncached_formulas: int = 0

    @property
    def is_uncalculated(self) -> bool:
        """Whether this column is formulas with no results behind them."""
        return self.uncached_formulas > 0 and self.populated == 0

    @property
    def needs_cast(self) -> str | None:
        """What this column really holds, when it is stored as text.

        Returned as a hint rather than applied here: this module reports facts,
        and rewriting a value is a decision for the extraction step.
        """
        if not self.populated:
            return None
        if self.date_as_text / self.populated >= 0.8:
            return "date"
        if self.numeric_as_text / self.populated >= 0.8:
            return "number"
        return None

    @property
    def dominant_kind(self) -> CellKind:
        if not self.kinds:
            return "empty"
        return max(self.kinds.items(), key=lambda kv: kv[1])[0]

    @property
    def dominant_unit(self) -> Unit:
        if not self.units:
            return "unknown"
        return max(self.units.items(), key=lambda kv: kv[1])[0]

    @property
    def dominant_currency(self) -> str | None:
        """The one currency this column is in, or None if it does not say.

        None both for a column that names none and for a column that MIXES
        them. Those differ upstream but are identical to a caller choosing a
        symbol: neither can be printed as one currency, and the honest
        rendering of both is no code rather than the majority one. A column
        that is 90% dirhams and 10% euros is a column whose total means
        nothing, and labelling it AED would hide exactly that.
        """
        if len(self.currencies) != 1:
            return None
        return next(iter(self.currencies))

    @property
    def mixes_currencies(self) -> bool:
        """More than one currency down one column -- worth saying out loud,
        because it makes every aggregate of that column meaningless."""
        return len(self.currencies) > 1

    @property
    def is_consistent(self) -> bool:
        """Whether the column holds one kind of thing all the way down. An
        inconsistent column usually means the region boundary is wrong."""
        if self.populated == 0:
            return False
        return self.kinds.get(self.dominant_kind, 0) / self.populated >= 0.8


@dataclass
class EntityField:
    """One label/value pair inside a form-shaped region."""

    label: str
    value: Any
    row: int
    unit: Unit
    is_derived: bool
    section: str | None = None
    #: The ISO code this field's own format names, if any. Per FIELD, not per
    #: region: on a form one line is dirhams and the line below it is a ratio,
    #: which is the whole reason this module reads formats cell by cell.
    currency: str | None = None


@dataclass
class Region:
    """A contiguous band of rows that holds one coherent thing."""

    kind: RegionKind
    first_row: int
    last_row: int
    first_col: int
    last_col: int

    # records
    header_row: int | None = None
    columns: list[ColumnProfile] = field(default_factory=list)
    data_first_row: int | None = None
    data_last_row: int | None = None

    # entity
    fields: list[EntityField] = field(default_factory=list)

    # both
    total_rows: list[int] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        if self.data_first_row is None or self.data_last_row is None:
            return 0
        return self.data_last_row - self.data_first_row + 1


def _row_bands(grid: Grid) -> list[tuple[int, int]]:
    """Split a sheet into bands of populated rows separated by blank rows.

    This is what makes two tables on one sheet tractable: the rent roll's
    residential table (rows 5-31) and commercial table (rows 33-39) are divided
    by the blank row 32, so they never get merged into one malformed region.
    """
    rows = grid.populated_rows()
    if not rows:
        return []

    bands: list[tuple[int, int]] = []
    start = prev = rows[0]
    for row in rows[1:]:
        if row - prev > REGION_GAP_ROWS:
            bands.append((start, prev))
            start = row
        prev = row
    bands.append((start, prev))
    return bands


def _has_total_label(grid: Grid, row: int) -> bool:
    """Whether a row calls itself a total."""
    for cell in grid.row_cells(row):
        if cell.kind == "text" and _TOTAL_LABEL.match(normalise_label(cell.value)):
            return True
    return False


def _is_all_computed(grid: Grid, row: int) -> bool:
    """Whether every figure in a row comes from a formula.

    A formula whose result was never saved counts too: it is still a computed
    cell, and skipping it would mean a subtotal row in an uncalculated workbook
    looked like ordinary data purely because its number is missing.
    """
    figures = [
        c
        for c in grid.row_cells(row)
        if c.kind == "number" or (c.kind == "empty" and c.is_derived)
    ]
    return bool(figures) and all(c.is_derived for c in figures)


def _is_total_row(grid: Grid, row: int) -> bool:
    """Whether a row aggregates the rows above it, judged in isolation.

    Used where there is no region to compare against — band splitting, and the
    orphan-total merge. Region-scoped callers should prefer `_total_rows_in`,
    which does not mistake a fully-computed sheet for a stack of subtotals.
    """
    return _has_total_label(grid, row) or _is_all_computed(grid, row)


def _total_rows_in(grid: Grid, rows: list[int]) -> list[int]:
    """Which of `rows` are aggregates of the others.

    "Every numeric cell is a formula" only means "subtotal" when the region ALSO
    contains rows that are not. On an operating statement the inputs are typed
    values and the subtotals are sums, so the signal is decisive. On a dashboard
    or exposure-analysis sheet EVERY figure is a cross-sheet formula, and
    applying the same rule marks every row a total and empties the table — the
    header is found, the data all discarded.

    So the formula signal is used only when it discriminates. When it does not,
    a row has to say "Total" to be treated as one.
    """
    labelled = [r for r in rows if _has_total_label(grid, r)]
    computed = [r for r in rows if _is_all_computed(grid, r)]

    if computed and len(computed) == len([r for r in rows if grid.row_cells(r)]):
        # Everything here is computed, so being computed says nothing.
        return sorted(set(labelled))
    return sorted(set(labelled) | set(computed))


def _next_content_row(grid: Grid, after: int, last_row: int) -> int | None:
    """The next row that carries data, skipping blanks and totals.

    Linear, so only for callers that ask once. Anything asking per row should
    build a `_ContentRows` index instead — see the note there.
    """
    for row in range(after + 1, last_row + 1):
        if not grid.is_blank_row(row) and not _is_total_row(grid, row):
            return row
    return None


class _ContentRows:
    """The rows of a band that carry data, indexed for repeated lookup.

    _find_header_row asks "what is the row after this one?" once per candidate.
    Answering by scanning forward is fine on a small band and catastrophic on a
    large computed one: nearly every row looks like a header candidate, and
    nearly every row looks like a total, so each scan runs thousands of rows
    deep before finding content. One real 10,650-row sheet turned that into
    ~100M checks and took thirteen minutes.

    Built once in a single pass, then answered by binary search.
    """

    def __init__(self, grid: Grid, first_row: int, last_row: int) -> None:
        self._rows = [
            row
            for row in range(first_row, last_row + 1)
            if not grid.is_blank_row(row) and not _is_total_row(grid, row)
        ]

    def after(self, row: int) -> int | None:
        index = bisect.bisect_right(self._rows, row)
        return self._rows[index] if index < len(self._rows) else None


def _is_unformatted(cell: Cell) -> bool:
    """Whether a cell carries no display formatting of its own.

    Headers are typed as plain labels and keep Excel's General format. Data
    cells in the same column routinely carry '$#,##0' or '0.0%'. When a data row
    is mostly text — tenant names, lease statuses, guaranty descriptions — this
    is the only thing left that separates it from the header above it.
    """
    return (cell.number_format or "General") == "General"


def _find_header_row(grid: Grid, first_row: int, last_row: int) -> int | None:
    """The first row in a band that names the columns beneath it.

    Two passes, because the obvious test is not sufficient. "Mostly text, and
    the row below is not mostly text" works for a table of numbers, but a rent
    roll's data rows are themselves mostly text, so that test picks the first
    tenant row as the header and discards the real one.

    The stronger signal is that a header is *entirely* text and *entirely*
    unformatted, while the row beneath it carries at least one number or one
    formatted cell. That resolves the rent roll correctly and still resolves
    everything the looser test used to.
    """
    candidates: list[tuple[int, float, bool]] = []
    for row in range(first_row, last_row + 1):
        cells = grid.row_cells(row)
        if len(cells) < 2:
            continue
        text = sum(1 for c in cells if c.kind == "text")
        ratio = text / len(cells)
        if ratio < HEADER_TEXT_RATIO:
            continue
        candidates.append((row, ratio, all(_is_unformatted(c) for c in cells)))

    if not candidates:
        return None

    content = _ContentRows(grid, first_row, last_row)

    def differs_from_row_below(row: int, strict: bool) -> bool:
        """Whether the row beneath `row` looks like data rather than more text.

        Returns False when there is nothing beneath — that is an absence of
        evidence, not evidence of a header, and treating it as decisive would
        make the LAST row of an all-text table win.
        """
        below = content.after(row)
        if below is None:
            return False
        below_cells = grid.row_cells(below)
        if not below_cells:
            return False
        if strict:
            return any(c.kind != "text" for c in below_cells) or any(
                not _is_unformatted(c) for c in below_cells
            )
        below_text = sum(1 for c in below_cells if c.kind == "text")
        return below_text / len(below_cells) < HEADER_TEXT_RATIO

    # Pass 1 — an all-text, unformatted row above a row that is neither.
    for row, ratio, unformatted in candidates:
        if ratio >= 1.0 and unformatted and differs_from_row_below(row, strict=True):
            return row

    # Pass 2 — the looser test, for headers that are not purely text.
    for row, ratio, _ in candidates:
        if differs_from_row_below(row, strict=False):
            return row

    # Nothing distinguished a header from its body: either a table whose every
    # column is text, or a header with a blank body beneath it. A header always
    # sits above its data, so the earliest candidate is the header.
    return candidates[0][0]


def _looks_like_entity(grid: Grid, first_row: int, last_row: int) -> bool:
    """Whether a band is a label/value form rather than a table.

    The test is column concentration: a form puts nearly everything in two
    columns (label, value), while a table spreads across many.
    """
    cols: dict[int, int] = {}
    total = 0
    for row in range(first_row, last_row + 1):
        for cell in grid.row_cells(row):
            cols[cell.col] = cols.get(cell.col, 0) + 1
            total += 1
    if total == 0:
        return False
    top_two = sorted(cols.values(), reverse=True)[:2]
    return sum(top_two) / total >= ENTITY_TWO_COL_RATIO


def _profile_records(grid: Grid, first_row: int, last_row: int, header_row: int) -> Region:
    header_cells = grid.row_cells(header_row)
    first_col = min(c.col for c in header_cells)
    last_col = max(c.col for c in header_cells)

    body = [
        r for r in range(header_row + 1, last_row + 1) if not grid.is_blank_row(r)
    ]
    total_rows = _total_rows_in(grid, body)
    # Membership against a SET, not the list. On a computed sheet most rows
    # qualify as totals, so `r not in total_rows` over a 10,000-row band is
    # 100M list scans — one real workbook spent thirteen minutes here.
    excluded = set(total_rows)
    data_rows = [r for r in body if r not in excluded]

    columns: list[ColumnProfile] = []
    for cell in header_cells:
        profile = ColumnProfile(col=cell.col, header=normalise_label(cell.value))
        for row in data_rows:
            data_cell = grid.at(row, cell.col)
            if data_cell is None:
                continue
            if data_cell.kind == "empty" and data_cell.is_derived:
                # A formula whose result was never stored. Counted, not treated
                # as a value — it is the difference between "this column is
                # empty" and "this column could not be read".
                profile.uncached_formulas += 1
                continue
            profile.populated += 1
            profile.kinds[data_cell.kind] = profile.kinds.get(data_cell.kind, 0) + 1
            unit = data_cell.unit
            profile.units[unit] = profile.units.get(unit, 0) + 1
            code = data_cell.currency
            if code:
                profile.currencies[code] = profile.currencies.get(code, 0) + 1
            if data_cell.is_derived:
                profile.derived_count += 1
            if data_cell.kind == "text":
                text = str(data_cell.value)
                if _DATE_TEXT.match(text):
                    profile.date_as_text += 1
                elif _NUMERIC_TEXT.match(text):
                    profile.numeric_as_text += 1
        columns.append(profile)

    region = Region(
        kind="records",
        first_row=first_row,
        last_row=last_row,
        first_col=first_col,
        last_col=last_col,
        header_row=header_row,
        columns=columns,
        data_first_row=min(data_rows) if data_rows else None,
        data_last_row=max(data_rows) if data_rows else None,
        total_rows=total_rows,
    )

    if not data_rows:
        region.notes.append(
            "Headers present with no data rows — a blank template. Not an error."
        )
    elif len(data_rows) < MIN_RECORD_ROWS:
        region.notes.append("Very few data rows; column types are weak evidence.")
    inconsistent = [c.header for c in columns if c.populated and not c.is_consistent]
    if inconsistent:
        region.notes.append(
            "Columns mixing types: " + ", ".join(inconsistent[:5])
        )
    casts = [f"{c.header} -> {c.needs_cast}" for c in columns if c.needs_cast]
    if casts:
        region.notes.append(
            "Stored as text but really typed (cast before aggregating): "
            + ", ".join(casts[:6])
        )
    uncalculated = [c.header for c in columns if c.is_uncalculated]
    if uncalculated:
        region.notes.append(
            "NOT CALCULATED — these columns are formulas whose results were "
            "never saved into the file, so their values cannot be read: "
            + ", ".join(uncalculated[:8])
            + ". Open the workbook in Excel and save it, or derive these from "
            "the source tables instead of reporting them as empty."
        )
    return region


def _profile_entity(grid: Grid, first_row: int, last_row: int) -> Region:
    """Extract a vertical form into label/value pairs, carrying section context.

    Section headings ("INCOME", "OPERATING EXPENSES") are all-caps rows with no
    value beside them. They are not data, but they qualify every field beneath
    them — "Insurance" means nothing until you know it sits under operating
    expenses.
    """
    fields: list[EntityField] = []
    total_rows: list[int] = []
    section: str | None = None

    aggregate_rows = set(
        _total_rows_in(
            grid,
            [r for r in range(first_row, last_row + 1) if not grid.is_blank_row(r)],
        )
    )

    populated_cols: dict[int, int] = {}
    for row in range(first_row, last_row + 1):
        for cell in grid.row_cells(row):
            populated_cols[cell.col] = populated_cols.get(cell.col, 0) + 1
    ordered = sorted(populated_cols, key=lambda c: (-populated_cols[c], c))
    label_col = min(ordered[:2]) if len(ordered) >= 2 else min(populated_cols, default=1)
    value_col = max(ordered[:2]) if len(ordered) >= 2 else label_col + 1

    for row in range(first_row, last_row + 1):
        label_cell = grid.at(row, label_col)
        value_cell = grid.at(row, value_col)
        if label_cell is None:
            continue
        label = normalise_label(label_cell.value)
        if not label:
            continue

        uncached = value_cell is not None and value_cell.kind == "empty" and value_cell.is_derived
        if value_cell is None or (value_cell.kind == "empty" and not uncached):
            # No value beside it: a section heading, or prose. A derived cell
            # whose value was never cached is NOT this case — it is a real
            # field we happen not to have the number for.
            if _SECTION_LABEL.match(label) and len(label) < 60:
                section = label
            continue

        # Region-scoped, for the same reason records regions are: on a form
        # where every line is computed, "this line is a formula" does not
        # single out the subtotals.
        is_total = row in aggregate_rows
        if is_total:
            total_rows.append(row)

        fields.append(
            EntityField(
                label=label,
                value=value_cell.value,
                row=row,
                unit=value_cell.unit,
                currency=value_cell.currency,
                is_derived=value_cell.is_derived or is_total,
                section=section,
            )
        )

    cols = [c for c in (label_col, value_col) if c]
    region = Region(
        kind="entity",
        first_row=first_row,
        last_row=last_row,
        first_col=min(cols),
        last_col=max(cols),
        fields=fields,
        total_rows=total_rows,
    )
    derived = sum(1 for f in fields if f.is_derived)
    if derived:
        region.notes.append(
            f"{derived} of {len(fields)} fields are derived — exclude them from sums."
        )
    uncached = sum(1 for f in fields if f.is_derived and f.value is None)
    if uncached:
        region.notes.append(
            f"{uncached} derived field(s) have no cached value — this workbook "
            "has never been calculated. The labels are reliable; the numbers "
            "must be recomputed from the inputs."
        )
    region.notes.append(
        "Read as a form because it occupies two columns. A narrow table with a "
        "header would look identical here — if the first row names the two "
        "columns rather than labelling a value, this is records, not fields."
    )
    return region


#: A lone text row spanning at least this many columns is a header, not prose.
#: A title or footnote occupies one or two cells; a header names every column.
MIN_HEADER_COLUMNS = 3


def _merge_orphan_totals(grid: Grid, regions: list[Region]) -> list[Region]:
    """Reattach a total row to the header it belongs to.

    A blank template separates its header from its total row by the empty rows
    the user is meant to fill in — the rent roll has a 17-column header on row 5,
    nothing until row 31, and "Total" there. Band splitting correctly sees two
    bands, but they are one table with zero rows in it.

    So: a records region with no data rows adopts the next region if that region
    is nothing but total rows and sits within its column span. The result
    reports the table's true extent, which is what a user confirming the mapping
    needs to see.
    """
    merged: list[Region] = []
    skip: set[int] = set()

    for index, region in enumerate(regions):
        if index in skip:
            continue
        if region.kind == "records" and region.row_count == 0:
            for offset in range(index + 1, len(regions)):
                candidate = regions[offset]
                rows = [
                    r
                    for r in range(candidate.first_row, candidate.last_row + 1)
                    if not grid.is_blank_row(r)
                ]
                if not rows or not all(_is_total_row(grid, r) for r in rows):
                    break
                overlaps = (
                    candidate.first_col <= region.last_col
                    and candidate.last_col >= region.first_col
                )
                if not overlaps:
                    break
                region.last_row = candidate.last_row
                region.total_rows = sorted(set(region.total_rows) | set(rows))
                region.notes.append(
                    f"Total row at {rows[0]} belongs to this table; the rows "
                    "between are the blank body of the template."
                )
                skip.add(offset)
                break
        merged.append(region)
    return merged


def profile_grid(grid: Grid) -> list[Region]:
    """Every distinct region on one sheet."""
    regions: list[Region] = []
    for first_row, last_row in _row_bands(grid):
        cells = [c for r in range(first_row, last_row + 1) for c in grid.row_cells(r)]
        if not cells:
            continue

        # Header detection runs FIRST. A header row is definitionally all text,
        # so testing "has no numbers, therefore prose" ahead of it discards the
        # headers of any table whose body is empty.
        header_row = _find_header_row(grid, first_row, last_row)
        is_entity = _looks_like_entity(grid, first_row, last_row)
        wide_enough = (
            header_row is not None
            and len(grid.row_cells(header_row)) >= MIN_HEADER_COLUMNS
        )
        numeric = [c for c in cells if c.kind in ("number", "date")]

        if header_row is not None and not is_entity and (numeric or wide_enough):
            regions.append(_profile_records(grid, first_row, last_row, header_row))
        elif is_entity and numeric:
            regions.append(_profile_entity(grid, first_row, last_row))
        elif not numeric:
            regions.append(
                Region(
                    kind="prose",
                    first_row=first_row,
                    last_row=last_row,
                    first_col=min(c.col for c in cells),
                    last_col=max(c.col for c in cells),
                )
            )
        else:
            regions.append(
                Region(
                    kind="unknown",
                    first_row=first_row,
                    last_row=last_row,
                    first_col=min(c.col for c in cells),
                    last_col=max(c.col for c in cells),
                    notes=["Neither a header/records table nor a label/value form."],
                )
            )
    return _merge_orphan_totals(grid, regions)


# ── workbook-level ────────────────────────────────────────────────────────────


@dataclass
class SheetProfile:
    name: str
    n_rows: int
    n_cols: int
    merged: list[str]
    regions: list[Region]
    #: The period this sheet says it covers, if it says so at all.
    as_of: date | None = None

    @property
    def kinds(self) -> list[RegionKind]:
        return [r.kind for r in self.regions]


@dataclass
class WorkbookProfile:
    path: str
    file_format: str
    sheets: list[SheetProfile]

    @property
    def as_of(self) -> date | None:
        """The period this workbook covers.

        Taken from the sheets that state one, falling back to the file name.
        Where sheets disagree — a workbook holding several months — the LATEST
        is used, because that is what the file as a whole is "as of". Per-sheet
        dates are still available for anything that needs finer grain.
        """
        stated = [s.as_of for s in self.sheets if s.as_of]
        if stated:
            return max(stated)
        return as_of_from_filename(Path(self.path).name)

    @property
    def fingerprint(self) -> str:
        """A hash of this workbook's SHAPE, ignoring its data.

        Built from the distinct region structures and their column headers, with
        sheet names and row counts left out. That makes it stable across the two
        things that always differ between files of the same kind — which
        property it covers, and how many rows it has — so next month's rent roll
        matches last month's and inherits the definitions already agreed for it.

        Freezing the layout is what a template does and why templates rot. This
        is only an index into saved meaning; when it misses, the mapper simply
        runs again.
        """
        parts: list[str] = []
        for signature, sheet_names in sorted(self.repeated_shapes.items()):
            exemplar = next(s for s in self.sheets if s.name == sheet_names[0])
            headers: list[str] = []
            for region in exemplar.regions:
                headers.extend(c.header.lower() for c in region.columns if c.header)
                headers.extend(f.label.lower() for f in region.fields)
            parts.append(f"{signature}::{'|'.join(sorted(headers))}")
        digest = hashlib.sha256("\n".join(parts).encode("utf-8")).hexdigest()
        return digest

    @property
    def repeated_shapes(self) -> dict[str, list[str]]:
        """Sheets grouped by identical region structure.

        The income statement's 20 property tabs are one shape repeated 20 times.
        Recognising that means the caller resolves meaning once and applies it
        twenty times — the same economy that makes a shape library worth having
        across uploads, happening here within a single file.
        """
        groups: dict[str, list[str]] = {}
        for sheet in self.sheets:
            signature = "|".join(
                f"{r.kind}:{r.last_col - r.first_col + 1}c" for r in sheet.regions
            )
            groups.setdefault(signature, []).append(sheet.name)
        return groups


#: Rows scanned for an "as of" line. It sits in the title block by convention;
#: scanning the whole sheet would find lease dates and call them the period.
_AS_OF_SCAN_ROWS = 12


def _sheet_as_of(grid: Grid) -> date | None:
    """The period a sheet announces in its heading, if any."""
    for row in range(1, min(_AS_OF_SCAN_ROWS, grid.n_rows) + 1):
        for cell in grid.row_cells(row):
            if cell.kind != "text":
                continue
            found = parse_as_of(str(cell.value))
            if found:
                return found
    return None


def profile_workbook(
    path: str | Path, grids: list[Grid] | None = None
) -> WorkbookProfile:
    """Describe every sheet of a workbook: where the data is and what is
    provably true about it.

    `grids` lets a caller that has already read the workbook hand them in.
    Reading is the expensive part — openpyxl parses the file twice, once for
    values and once for formulas — so a pipeline that profiles and then loads
    the same file should pay for that once, not twice.
    """
    path = Path(path)
    if grids is None:
        grids = read_grids(path)
    sheets = [
        SheetProfile(
            name=g.name,
            n_rows=g.n_rows,
            n_cols=g.n_cols,
            merged=g.merged,
            regions=profile_grid(g),
            as_of=_sheet_as_of(g),
        )
        for g in grids
    ]
    return WorkbookProfile(
        path=str(path), file_format=path.suffix.lower().lstrip("."), sheets=sheets
    )


# ── evidence packet ───────────────────────────────────────────────────────────


def evidence_packet(profile: WorkbookProfile, sample_rows: int = 3) -> dict[str, Any]:
    """A compact, JSON-safe description of the workbook, sized for a prompt.

    This is the seam between the deterministic half of the system and the
    semantic half. Everything in here is a fact we proved; nothing in here is a
    guess about meaning. The caller hands this to a model and asks only the
    question a model should answer: what do these fields mean?

    One entry per DISTINCT shape rather than per sheet, so a 20-tab workbook of
    identical statements costs one shape's worth of context, not twenty.
    """
    by_shape = profile.repeated_shapes
    seen: set[str] = set()

    shapes: list[dict[str, Any]] = []
    for signature, sheet_names in by_shape.items():
        if signature in seen:
            continue
        seen.add(signature)
        exemplar = next(s for s in profile.sheets if s.name == sheet_names[0])

        regions: list[dict[str, Any]] = []
        for region in exemplar.regions:
            entry: dict[str, Any] = {
                "kind": region.kind,
                "rows": f"{region.first_row}-{region.last_row}",
                "notes": region.notes,
            }
            if region.kind == "records":
                entry["header_row"] = region.header_row
                entry["data_rows"] = region.row_count
                entry["total_rows"] = region.total_rows
                entry["columns"] = [
                    {
                        "col": c.col,
                        "header": c.header,
                        "type": c.dominant_kind,
                        "unit": c.dominant_unit,
                        # Only present when the workbook actually named one.
                        # An absent key is "not stated", which downstream
                        # renders as dollars; an empty string is "this column
                        # mixes currencies", which must NOT render as dollars.
                        # Collapsing the two is the bug this whole path exists
                        # to close, so they stay separate all the way out.
                        **({"currency": c.dominant_currency}
                           if c.dominant_currency else {}),
                        **({"currency": "",
                            "currency_note": "mixes "
                            + ", ".join(sorted(c.currencies))}
                           if c.mixes_currencies else {}),
                        "consistent": c.is_consistent,
                        "derived": c.derived_count > 0,
                    }
                    for c in region.columns
                ]
            elif region.kind == "entity":
                entry["fields"] = [
                    {
                        "label": f.label,
                        "section": f.section,
                        "unit": f.unit,
                        **({"currency": f.currency} if f.currency else {}),
                        "derived": f.is_derived,
                        "sample": _jsonable(f.value),
                    }
                    for f in region.fields[: sample_rows * 10]
                ]
            regions.append(entry)

        shapes.append(
            {
                "sheets": sheet_names,
                "sheet_count": len(sheet_names),
                "regions": regions,
            }
        )

    return {
        "file": Path(profile.path).name,
        "format": profile.file_format,
        "sheet_count": len(profile.sheets),
        "distinct_shapes": len(shapes),
        "shapes": shapes,
    }


def _jsonable(value: Any) -> Any:
    import datetime as _dt

    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def iter_all_regions(profile: WorkbookProfile) -> Iterable[tuple[str, Region]]:
    for sheet in profile.sheets:
        for region in sheet.regions:
            yield sheet.name, region
