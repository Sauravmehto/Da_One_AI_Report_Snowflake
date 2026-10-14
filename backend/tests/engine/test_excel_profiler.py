"""Profiler tests, built from the two shapes real client workbooks actually take.

Every case here is a reduction of something observed in a real file, not an
invented edge case:

  * `entity` mirrors a 20-tab operating statement — line items down one column,
    values beside them, section headings interleaved, and subtotals sitting in
    the middle of the data rather than at the end.
  * `records` mirrors a rent-roll template — two tables on one sheet whose
    columns do not align, a header separated from its total row by the empty
    body a user is meant to fill in, and headers carrying newlines.

The subtotal cases are deliberately split: the operating statement marks its
totals with formulas, the rent roll marks its with the word "Total" and a
literal zero. Both must be caught, which is why the profiler tests for each.
"""

from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from app.engine.excel_profiler import (
    Cell,
    Grid,
    _total_rows_in,
    as_of_from_filename,
    normalise_label,
    parse_as_of,
    profile_grid,
    profile_workbook,
    unit_from_format,
)


# ── fixtures ──────────────────────────────────────────────────────────────────


@pytest.fixture
def statement_workbook(tmp_path):
    """A property operating statement: a form, not a table."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01 Some Property"

    ws["A1"] = "Some Property"
    ws["A2"] = "123 Example St"

    ws["A4"] = "INCOME"
    ws["A5"], ws["B5"] = "Gross potential rent", 1_000_000
    ws["A6"], ws["B6"] = "Other income", 50_000
    ws["A7"], ws["B7"] = "Effective gross income (EGI)", "=SUM(B5:B6)"

    ws["A9"] = "OPERATING EXPENSES"
    ws["A10"], ws["B10"] = "Real estate taxes", 80_000
    ws["A11"], ws["B11"] = "Insurance", 20_000
    ws["A12"], ws["B12"] = "Total operating expenses", "=SUM(B10:B11)"

    ws["A14"], ws["B14"] = "Loan-to-value (LTV)", 0.625
    ws["B14"].number_format = "0.0%"
    for coordinate in ("B5", "B6", "B7", "B10", "B11", "B12"):
        ws[coordinate].number_format = "$#,##0"

    path = tmp_path / "statement.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def rent_roll_workbook(tmp_path):
    """A blank rent-roll template: two tables, no data, totals far below."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rent Roll Template"

    ws["B1"] = "Rent Roll Schedule"
    ws["B2"] = "[Property Name]"

    headers = [
        "Street Address",
        "Unit #",
        "Unit Square footage",
        "Lease Rent",
        "Comments\n(Add notes as needed)",
    ]
    for offset, text in enumerate(headers):
        ws.cell(row=4, column=2 + offset, value=text)
    # Rows 5-14 are the blank body a user fills in.
    ws["B15"] = "Total"
    ws["E15"] = 0.0

    for offset, text in enumerate(["Commercial Tenant Name", "Unit Square footage", "Lease Rent"]):
        ws.cell(row=17, column=2 + offset, value=text)
    ws["B21"] = "Total"
    ws["D21"] = 0.0

    path = tmp_path / "rent_roll.xlsx"
    wb.save(path)
    return path


# ── units ─────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "number_format,expected",
    [
        ('\\$#,##0;"($"#,##0\\);\\-', "currency"),
        ("0.0%", "percent"),
        ("#,##0", "number"),
        ("yyyy-mm-dd", "date"),
        (None, "number"),
    ],
)
def test_the_number_format_decides_the_unit(number_format, expected):
    """One column can hold dollars and ratios at once, so the unit has to come
    from each cell's own format rather than from the column."""
    assert unit_from_format(number_format, "number") == expected


# ── the period a file covers ──────────────────────────────────────────────────


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Rent Roll Detail — as of June 30, 2026", date(2026, 6, 30)),
        ("As Of: 06/30/2026", date(2026, 6, 30)),
        ("For the quarter ended March 31, 2026", date(2026, 3, 31)),
        ("Period ending 2026-06-30", date(2026, 6, 30)),
        # A bare date is NOT a period. Reading a lease expiry as the file's
        # as-of would file a whole month's rows under the wrong month.
        ("Lease End 06/30/2026", None),
        ("Printed 12/30/2023", None),
        ("Bayou Crossing Industrial Park", None),
    ],
)
def test_only_an_announced_period_counts_as_the_as_of(text, expected):
    assert parse_as_of(text) == expected


@pytest.mark.parametrize(
    "filename,expected",
    [
        ("RentRoll_2026-06.xlsx", date(2026, 6, 1)),
        ("Rent Roll Jun 2026.xlsx", date(2026, 6, 1)),
        ("rent_roll_20260630.xlsx", date(2026, 6, 30)),
        ("Rent Roll Final.xlsx", None),
    ],
)
def test_the_filename_is_a_fallback_when_the_sheet_is_silent(filename, expected):
    assert as_of_from_filename(filename) == expected


def test_a_workbook_reports_the_period_its_sheets_announce(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Portfolio Rent Roll — as of March 31, 2026"
    for offset, text in enumerate(["Unit", "Tenant", "Rent"]):
        ws.cell(row=3, column=1 + offset, value=text)
    ws["A4"], ws["B4"], ws["C4"] = "101", "Ana Diaz", 2400
    path = tmp_path / "no_date_in_name.xlsx"
    wb.save(path)

    assert profile_workbook(path).as_of == date(2026, 3, 31)


def test_a_header_with_a_newline_collapses_to_one_line():
    assert normalise_label("Comments\n(Add notes as needed)") == (
        "Comments (Add notes as needed)"
    )


# ── entity shape ──────────────────────────────────────────────────────────────


def test_a_statement_is_read_as_fields_not_records(statement_workbook):
    profile = profile_workbook(statement_workbook)
    kinds = [r.kind for r in profile.sheets[0].regions]
    assert "entity" in kinds
    assert "records" not in kinds


def test_section_headings_qualify_the_fields_beneath_them(statement_workbook):
    profile = profile_workbook(statement_workbook)
    fields = [f for r in profile.sheets[0].regions for f in r.fields]
    by_label = {f.label: f for f in fields}

    assert by_label["Gross potential rent"].section == "INCOME"
    assert by_label["Insurance"].section == "OPERATING EXPENSES"


def test_formula_subtotals_are_flagged_so_sums_do_not_double_count(statement_workbook):
    profile = profile_workbook(statement_workbook)
    by_label = {
        f.label: f for r in profile.sheets[0].regions for f in r.fields
    }

    assert by_label["Effective gross income (EGI)"].is_derived
    assert by_label["Total operating expenses"].is_derived
    assert not by_label["Gross potential rent"].is_derived
    assert not by_label["Insurance"].is_derived


def test_dollars_and_percentages_in_one_column_keep_their_own_units(
    statement_workbook,
):
    profile = profile_workbook(statement_workbook)
    by_label = {
        f.label: f for r in profile.sheets[0].regions for f in r.fields
    }

    assert by_label["Gross potential rent"].unit == "currency"
    assert by_label["Loan-to-value (LTV)"].unit == "percent"


# ── records shape ─────────────────────────────────────────────────────────────


def test_two_tables_on_one_sheet_are_found_separately(rent_roll_workbook):
    profile = profile_workbook(rent_roll_workbook)
    records = [r for r in profile.sheets[0].regions if r.kind == "records"]

    assert len(records) == 2
    assert records[0].header_row == 4
    assert records[1].header_row == 17


def test_a_blank_template_keeps_its_headers(rent_roll_workbook):
    """The header row of an empty table has no numbers in it. Concluding
    "no numbers, therefore prose" would discard the entire schema."""
    profile = profile_workbook(rent_roll_workbook)
    first = [r for r in profile.sheets[0].regions if r.kind == "records"][0]

    assert [c.header for c in first.columns] == [
        "Street Address",
        "Unit #",
        "Unit Square footage",
        "Lease Rent",
        "Comments (Add notes as needed)",
    ]
    assert first.row_count == 0
    assert any("blank template" in n for n in first.notes)


def test_a_total_row_below_the_blank_body_rejoins_its_table(rent_roll_workbook):
    profile = profile_workbook(rent_roll_workbook)
    first = [r for r in profile.sheets[0].regions if r.kind == "records"][0]

    assert 15 in first.total_rows
    assert first.last_row == 15


def test_a_literal_total_is_caught_without_a_formula(rent_roll_workbook):
    """The rent roll's total row holds a literal 0.0, so the formula signal
    never fires and only the label identifies it."""
    profile = profile_workbook(rent_roll_workbook)
    records = [r for r in profile.sheets[0].regions if r.kind == "records"]

    assert all(r.total_rows for r in records)


def test_the_title_block_is_not_mistaken_for_a_header(rent_roll_workbook):
    profile = profile_workbook(rent_roll_workbook)
    prose = [r for r in profile.sheets[0].regions if r.kind == "prose"]

    assert any(r.first_row == 1 for r in prose)


# ── headers above text-heavy data ─────────────────────────────────────────────


@pytest.fixture
def tenant_workbook(tmp_path):
    """A tenant rent roll, where the DATA rows are themselves mostly text.

    Tenant name, lease status, guarantor, credit tier and payment history are
    all text, so "the row below is not mostly text" never fires and the first
    tenant row gets mistaken for the header. What separates them is that the
    header is entirely text and entirely unformatted, while the data row below
    carries currency formats.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01 Some Industrial Park"

    headers = ["Suite", "Type", "SF", "Tenant", "Rent/Mo", "Lease Start", "Status", "Guarantor"]
    for offset, text in enumerate(headers):
        ws.cell(row=3, column=1 + offset, value=text)

    rows = [
        ("Suite 1", "Warehouse", 142000, "Legacy Packaging LLC", 113990, "12/30/2023", "Current", "Corporate guaranty"),
        ("Suite 2", "Warehouse", 96000, "Harbor Freight Co", 74500, "03/15/2024", "Current", "Personal guaranty"),
    ]
    for row_offset, values in enumerate(rows):
        for col_offset, value in enumerate(values):
            ws.cell(row=4 + row_offset, column=1 + col_offset, value=value)
        ws.cell(row=4 + row_offset, column=5).number_format = "$#,##0"

    path = tmp_path / "tenants.xlsx"
    wb.save(path)
    return path


def test_a_text_heavy_data_row_is_not_mistaken_for_the_header(tenant_workbook):
    profile = profile_workbook(tenant_workbook)
    records = [r for r in profile.sheets[0].regions if r.kind == "records"][0]

    assert records.header_row == 3
    assert [c.header for c in records.columns][:4] == ["Suite", "Type", "SF", "Tenant"]
    assert records.row_count == 2


def test_currency_formatting_survives_into_the_column_profile(tenant_workbook):
    profile = profile_workbook(tenant_workbook)
    records = [r for r in profile.sheets[0].regions if r.kind == "records"][0]
    by_header = {c.header: c for c in records.columns}

    assert by_header["Rent/Mo"].dominant_unit == "currency"
    assert by_header["SF"].dominant_unit == "number"


def test_dates_stored_as_text_are_reported_as_needing_a_cast(tenant_workbook):
    """Rent rolls store dates as strings constantly. The column is text, and
    saying so without saying what it really holds would be useless."""
    profile = profile_workbook(tenant_workbook)
    records = [r for r in profile.sheets[0].regions if r.kind == "records"][0]
    by_header = {c.header: c for c in records.columns}

    assert by_header["Lease Start"].dominant_kind == "text"
    assert by_header["Lease Start"].needs_cast == "date"
    assert by_header["Tenant"].needs_cast is None


def _grid(rows: dict[int, list[tuple[object, str | None]]]) -> Grid:
    """A Grid built directly from (value, formula) pairs.

    Constructed rather than saved-and-reloaded because openpyxl cannot write a
    CACHED formula result — Excel stores the last computed value, openpyxl
    stores only the expression. Testing the subtotal rules through a file would
    therefore only ever exercise the uncalculated case.
    """
    grid = Grid(name="s", n_rows=max(rows), n_cols=3)
    for row, cells in rows.items():
        for index, (value, formula) in enumerate(cells, start=1):
            kind = "empty" if value is None else ("text" if isinstance(value, str) else "number")
            grid.cells[(row, index)] = Cell(
                row=row, col=index, value=value, kind=kind, formula=formula
            )
    return grid


def test_a_fully_computed_region_is_not_read_as_all_subtotals():
    """A dashboard sheet is formulas end to end.

    "Every figure is a formula" identifies a subtotal only when the region also
    holds rows that are not. On an exposure-analysis tab where every number is a
    cross-sheet formula, applying that rule marked every row a total and left
    the table with a header and no data — exactly the reported symptom: the
    table was detected but had nothing in it.
    """
    grid = _grid(
        {
            2: [("Software", None), (4, "=COUNTIF(x)"), (1000, "=SUMIF(x)")],
            3: [("Healthcare", None), (6, "=COUNTIF(y)"), (2000, "=SUMIF(y)")],
            4: [("Industrials", None), (2, "=COUNTIF(z)"), (500, "=SUMIF(z)")],
        }
    )
    assert _total_rows_in(grid, [2, 3, 4]) == []


def test_a_real_subtotal_is_still_caught_when_the_region_is_mixed():
    """The formula signal must keep working where it discriminates."""
    grid = _grid(
        {
            2: [("Software", None), (4, None), (1000, None)],
            3: [("Healthcare", None), (6, None), (2000, None)],
            4: [("All sectors", None), (10, "=SUM(B2:B3)"), (3000, "=SUM(C2:C3)")],
        }
    )
    assert _total_rows_in(grid, [2, 3, 4]) == [4]


def test_a_large_computed_sheet_profiles_in_reasonable_time():
    """Guards a quadratic that made one real workbook take thirteen minutes.

    The shape that triggers it: a big sheet where nearly every row looks like a
    header candidate (mostly text) AND nearly every row looks like a total
    (every figure computed). Header detection then asked "what is the next
    content row?" once per candidate, and each answer scanned thousands of rows
    forward before finding one, so the work grew with the square of the sheet.

    Timed rather than counted because the failure is a stall, not a wrong
    answer — asserting on structure would have passed throughout.
    """
    import time

    rows: dict[int, list[tuple[object, str | None]]] = {}
    for row in range(1, 4001):
        rows[row] = [
            (f"Lien {row}", None),
            (f"County {row % 50}", None),
            (float(row), f"=SUM(Raw!A{row})"),
        ]
    grid = _grid(rows)

    started = time.perf_counter()
    regions = profile_grid(grid)
    elapsed = time.perf_counter() - started

    assert regions, "the sheet should still yield a region"
    # Comfortably above the ~0.2s this takes, low enough that a return of the
    # quadratic (which ran to minutes) fails immediately.
    assert elapsed < 5.0, f"profiling 4,000 rows took {elapsed:.1f}s"


def test_a_row_calling_itself_a_total_is_caught_either_way():
    """The label is decisive on its own — the rent roll's total row holds a
    literal 0.0, so no formula signal exists to find it by."""
    grid = _grid(
        {
            2: [("Unit 1", None), (900, None), (2400, None)],
            3: [("Total", None), (900, None), (2400, None)],
        }
    )
    assert _total_rows_in(grid, [2, 3]) == [3]


def test_formulas_with_no_saved_results_are_reported_not_silently_empty(tmp_path):
    """A workbook Excel has never calculated carries formulas but no values.

    openpyxl cannot evaluate them, so the column loads as entirely NULL. Saying
    nothing would leave the model reporting "no exposure" for a file whose
    numbers simply were not saved — confidently wrong rather than visibly
    missing.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    for offset, text in enumerate(["Sector", "Region", "Principal"]):
        ws.cell(row=1, column=1 + offset, value=text)
    ws["A2"], ws["B2"], ws["C2"] = "Software", "EMEA", "=SUM(LOANS!C:C)"
    ws["A3"], ws["B3"], ws["C3"] = "Healthcare", "APAC", "=SUM(LOANS!D:D)"
    path = tmp_path / "uncalculated.xlsx"
    wb.save(path)

    region = [
        r for r in profile_workbook(path).sheets[0].regions if r.kind == "records"
    ][0]
    by_header = {c.header: c for c in region.columns}

    assert by_header["Principal"].is_uncalculated
    assert not by_header["Sector"].is_uncalculated
    assert any("NOT CALCULATED" in note for note in region.notes)


def test_numbers_stored_as_text_are_reported_as_needing_a_cast(tmp_path):
    """A rent-roll summary row stores '100%' and '$113,990' as strings."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "Property", "Occupancy", "Monthly Rent"
    ws["A2"], ws["B2"], ws["C2"] = "Cortez Point", "100%", "$113,990"
    ws["A3"], ws["B3"], ws["C3"] = "Rio Ridge", "94%", "$88,400"
    path = tmp_path / "text_numbers.xlsx"
    wb.save(path)

    profile = profile_workbook(path)
    records = [r for r in profile.sheets[0].regions if r.kind == "records"][0]
    by_header = {c.header: c for c in records.columns}

    assert by_header["Monthly Rent"].needs_cast == "number"
    assert by_header["Occupancy"].needs_cast == "number"
    assert by_header["Property"].needs_cast is None


def test_a_two_column_region_admits_it_is_ambiguous(tmp_path):
    """Two columns of label-and-value could be a form or a narrow table, and
    nothing in the structure settles it. Guessing silently would be worse than
    saying so — the caller resolving meaning can decide."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Real estate taxes", 80_000
    ws["A2"], ws["B2"] = "Insurance", 20_000
    ws["A3"], ws["B3"] = "Utilities", 9_000
    path = tmp_path / "two_col.xlsx"
    wb.save(path)

    region = profile_workbook(path).sheets[0].regions[0]

    assert region.kind == "entity"
    assert any("two columns" in n.lower() for n in region.notes)


# ── workbook level ────────────────────────────────────────────────────────────


def test_identical_sheets_collapse_into_one_shape(tmp_path):
    """A 20-tab workbook of identical statements should cost one shape's worth
    of attention, not twenty."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for index in range(5):
        ws = wb.create_sheet(f"{index:02d} Property")
        ws["A1"] = "INCOME"
        ws["A2"], ws["B2"] = "Gross potential rent", 1000 * (index + 1)
        ws["A3"], ws["B3"] = "Insurance", 100 * (index + 1)
    path = tmp_path / "many.xlsx"
    wb.save(path)

    profile = profile_workbook(path)
    assert len(profile.sheets) == 5
    assert len(profile.repeated_shapes) == 1


def test_the_evidence_packet_describes_shapes_not_every_sheet(tmp_path):
    from app.engine.excel_profiler import evidence_packet

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for index in range(8):
        ws = wb.create_sheet(f"{index:02d} Property")
        ws["A1"] = "INCOME"
        ws["A2"], ws["B2"] = "Gross potential rent", 1000
    path = tmp_path / "many.xlsx"
    wb.save(path)

    packet = evidence_packet(profile_workbook(path))
    assert packet["sheet_count"] == 8
    assert packet["distinct_shapes"] == 1
    assert packet["shapes"][0]["sheet_count"] == 8


def test_an_unsupported_format_is_refused_by_name(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported workbook format"):
        profile_workbook(path)
