"""Dataset identity and naming.

A dataset is a collection somebody uploaded, not a shape. These cover the rules
that keep two collections built on one template apart, and keep a dataset's name
meaning something once it holds more than one file.
"""

from __future__ import annotations

from pathlib import Path

import duckdb
import pytest

from app.engine.excel_materialize import MaterializedTable
from app.engine.excel_sources import (
    Dataset,
    SourceTable,
    _folder_of,
    _merge_tables,
    _name_for,
    _prune_stale_tables,
    _Prepared,
    _resolve_name,
    _unique,
    new_dataset_id,
)


def _prepared(filename: str, folder: str = "") -> _Prepared:
    return _Prepared(
        filename=filename, data=b"", path=Path(filename), profile=None, grids=[],
        fingerprint="f" * 64, as_of=None, as_of_source="", folder=folder,
    )


def _table(name: str, rows: int = 10, columns: list[str] | None = None):
    return MaterializedTable(
        name=name, row_count=rows, columns=columns or ["a"], sheets=["S"]
    )


# -- identity -----------------------------------------------------------------


def test_two_uploads_of_one_shape_get_different_ids():
    """Shape must not decide identity, or a second client's files join the first's."""
    assert new_dataset_id() != new_dataset_id()


def test_a_dataset_id_is_a_short_hex_string():
    identity = new_dataset_id()
    assert len(identity) == 16
    int(identity, 16)


# -- where a name comes from --------------------------------------------------


@pytest.mark.parametrize(
    "relative_path,expected",
    [
        ("Aug2026 Deals/Commercial.xlsx", "Aug2026 Deals"),
        ("Rent Rolls/2026/January.xlsx", "Rent Rolls"),
        ("Commercial.xlsx", ""),
    ],
)
def test_the_folder_is_read_from_the_picked_path(relative_path, expected):
    assert _folder_of(relative_path) == expected


def test_the_folder_names_the_dataset():
    """What the user picked, not whichever file happened to sort first."""
    items = [
        _prepared("Commercial.xlsx", folder="Aug 2026 Deals"),
        _prepared("Land.xlsx", folder="Aug 2026 Deals"),
    ]
    assert _name_for([_table("deal_pipeline")], items) == "Aug 2026 Deals"


def test_files_picked_without_a_folder_are_named_from_their_table():
    items = [_prepared("whatever.xlsx")]
    assert _name_for([_table("deal_pipeline")], items) == "Deal Pipeline"


def test_an_anonymous_table_does_not_name_a_dataset():
    """A digest is an identifier, not something to show a user."""
    items = [_prepared("Q3 Deals.xlsx")]
    assert _name_for([_table("shape_deal_id_city_a1b2c3")], items) == "Q3 Deals"


def test_a_folder_that_fans_out_does_not_name_every_dataset_after_itself():
    """Four shapes in one folder cannot all be called after the folder."""
    items = [_prepared("rent_roll.xlsx", folder="Mixed Folder")]
    assert _name_for([_table("unit_rent_roll")], items, use_folder=False) == (
        "Unit Rent Roll"
    )
    anonymous = [_table("shape_unit_tenant_sf_36c50d")]
    assert _name_for(anonymous, items, use_folder=False) == "rent roll"


def test_files_from_different_folders_fall_back_to_the_table():
    items = [
        _prepared("a.xlsx", folder="One"),
        _prepared("b.xlsx", folder="Two"),
    ]
    assert _name_for([_table("deal_pipeline")], items) == "Deal Pipeline"


# -- keeping the library legible ----------------------------------------------


def test_a_repeated_name_is_made_distinguishable():
    taken = {"Deal Pipeline"}
    assert _unique("Deal Pipeline", taken) == "Deal Pipeline (2)"
    assert _unique("Deal Pipeline", taken | {"Deal Pipeline (2)"}) == "Deal Pipeline (3)"


def test_a_free_name_is_left_alone():
    assert _unique("Rent Roll", {"Deal Pipeline"}) == "Rent Roll"


def test_adding_files_does_not_rename_a_dataset():
    """More rent rolls arriving is not a reason to relabel the set."""
    existing = Dataset(
        dataset_id="abc", name="Aug 2026 Deals", fingerprint="f",
        tables=[SourceTable(name="deal_pipeline", row_count=5, columns=[], sheets=[])],
    )
    name = _resolve_name(
        existing,
        [_table("deal_pipeline")],
        [_prepared("March.xlsx", folder="Something Else")],
        {"Aug 2026 Deals"},
    )
    assert name == "Aug 2026 Deals"


def test_a_new_dataset_of_a_shape_already_in_the_library_is_named_apart():
    """Two folders on one template must be distinguishable in the list."""
    items = [_prepared("a.xlsx", folder="Deals")]
    assert _resolve_name(None, [_table("deal_pipeline")], items, {"Deals"}) == "Deals (2)"


# -- merging what each file reported ------------------------------------------


def test_columns_and_unreadable_warnings_accumulate_across_files():
    into: dict = {}
    first = _table("deals", columns=["deal_id", "close_price"])
    first.unreadable_columns = ["close_price"]
    first.uncalculated_columns = ["commission", "days_to_close"]
    second = _table("deals", columns=["deal_id", "close_price", "notes"])
    second.unreadable_columns = []
    second.uncalculated_columns = ["commission"]

    _merge_tables(into, [first])
    _merge_tables(into, [second])

    merged = into["deals"]
    assert merged.columns == ["deal_id", "close_price", "notes"]
    # Unreadable in any file means a total over it is wrong.
    assert merged.unreadable_columns == ["close_price"]
    # A formula result one file saved is not missing from the dataset.
    assert merged.uncalculated_columns == ["commission"]


# -- clearing what an earlier naming scheme left behind -----------------------


def test_tables_from_an_earlier_naming_scheme_are_dropped(tmp_path):
    database = tmp_path / "d.duckdb"
    conn = duckdb.connect(str(database))
    for name in ("residential_sales_4_22", "deal_pipeline", "deal_pipeline_derived"):
        conn.execute(f'CREATE TABLE "{name}" (x INTEGER)')
    conn.close()

    _prune_stale_tables(database, {"deal_pipeline"})

    conn = duckdb.connect(str(database), read_only=True)
    try:
        remaining = {row[0] for row in conn.execute("SHOW TABLES").fetchall()}
    finally:
        conn.close()
    # The companion is kept: it is written without being reported as a table.
    assert remaining == {"deal_pipeline", "deal_pipeline_derived"}


def test_nothing_is_dropped_when_a_batch_produced_nothing(tmp_path):
    database = tmp_path / "d.duckdb"
    conn = duckdb.connect(str(database))
    conn.execute('CREATE TABLE "keep_me" (x INTEGER)')
    conn.close()

    _prune_stale_tables(database, set())

    conn = duckdb.connect(str(database), read_only=True)
    try:
        assert {row[0] for row in conn.execute("SHOW TABLES").fetchall()} == {"keep_me"}
    finally:
        conn.close()


# -- taking one file back out -------------------------------------------------


@pytest.fixture
def loaded(tmp_path, monkeypatch):
    """A dataset of three same-shaped files, with persistence held in memory."""
    import openpyxl

    from app.engine import excel_sources as module

    saved: dict[str, Dataset] = {}
    store_dir = tmp_path / "store"
    store_dir.mkdir()

    class Store:
        def put(self, key, data):
            (store_dir / key.replace("/", "__")).write_bytes(data)

        def get(self, key):
            return (store_dir / key.replace("/", "__")).read_bytes()

    monkeypatch.setattr(module, "_ensure_tables", lambda: True)
    monkeypatch.setattr(module, "_save_dataset", lambda d: saved.__setitem__(d.dataset_id, d))
    monkeypatch.setattr(module, "list_datasets", lambda: list(saved.values()))
    monkeypatch.setattr(module, "get_dataset", lambda i: saved.get(i))
    monkeypatch.setattr(module, "load_definition", lambda fp: None)
    monkeypatch.setattr(module, "_save_definition", lambda *a, **k: None)
    monkeypatch.setattr(module, "known_concepts", lambda: [])
    monkeypatch.setattr(module, "delete_dataset", lambda i: saved.pop(i, None) is not None)
    monkeypatch.setattr(module.storage_mod, "get_storage", lambda: Store())

    payload = []
    for name, count in (("a.xlsx", 3), ("b.xlsx", 5), ("c.xlsx", 2)):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(name.split(".")[0])
        ws["A1"] = "Deals"
        for i, header in enumerate(["Deal ID", "City", "Close Price"]):
            ws.cell(row=3, column=1 + i, value=header)
        for r in range(count):
            ws.cell(row=4 + r, column=1, value=f"{name}-{r}")
            ws.cell(row=4 + r, column=2, value="Newark")
            ws.cell(row=4 + r, column=3, value=100 + r)
        path = tmp_path / name
        wb.save(path)
        payload.append((f"Deals/{name}", path.read_bytes()))

    result = module.ingest_files(payload, owner="me")
    return module, saved, result.datasets[0]


def test_removing_a_file_deletes_only_its_rows(loaded):
    module, saved, dataset = loaded
    assert dataset.file_count == 3
    before = dataset.total_rows

    outcome = module.remove_file(dataset.dataset_id, "b.xlsx")

    assert outcome.removed and not outcome.dataset_deleted
    after = saved[dataset.dataset_id]
    assert after.file_count == 2
    assert [f.filename for f in after.files] == ["a.xlsx", "c.xlsx"]
    assert after.total_rows == before - 5


def test_removing_a_file_that_is_not_there_changes_nothing(loaded):
    module, saved, dataset = loaded
    assert module.remove_file(dataset.dataset_id, "nope.xlsx").removed is False
    assert saved[dataset.dataset_id].file_count == 3


def test_removing_the_last_file_removes_the_dataset(loaded):
    """A dataset with no files answers every question with nothing."""
    module, saved, dataset = loaded
    for name in ("a.xlsx", "b.xlsx"):
        assert module.remove_file(dataset.dataset_id, name).removed

    outcome = module.remove_file(dataset.dataset_id, "c.xlsx")

    assert outcome.dataset_deleted
    assert dataset.dataset_id not in saved


def test_the_folder_is_recorded_on_every_file(loaded):
    _, _, dataset = loaded
    assert dataset.folder == "Deals"
    assert all(f.folder == "Deals" for f in dataset.files)


# -- the same folder, uploaded twice ------------------------------------------


def test_uploading_one_folder_twice_gives_two_distinguishable_folders(
    tmp_path, monkeypatch
):
    """People re-upload the same folder, deliberately or by accident.

    Each upload is its own collection by design — two folders on one template
    are different data. But two rows reading "DA EXPRESS" with identical counts
    leave no way to tell which is which, or which one to delete.
    """
    import openpyxl

    from app.engine import excel_sources as module

    saved: dict[str, Dataset] = {}
    names: dict[str, str] = {}
    store_dir = tmp_path / "store"
    store_dir.mkdir()

    class Store:
        def put(self, key, data):
            (store_dir / key.replace("/", "__")).write_bytes(data)

        def get(self, key):
            return (store_dir / key.replace("/", "__")).read_bytes()

    monkeypatch.setattr(module, "_ensure_tables", lambda: True)
    monkeypatch.setattr(module, "_save_dataset", lambda d: saved.__setitem__(d.dataset_id, d))
    monkeypatch.setattr(module, "list_datasets", lambda: list(saved.values()))
    monkeypatch.setattr(module, "get_dataset", lambda i: saved.get(i))
    monkeypatch.setattr(module, "load_definition", lambda fp: None)
    monkeypatch.setattr(module, "_save_definition", lambda *a, **k: None)
    monkeypatch.setattr(module, "known_concepts", lambda: [])
    monkeypatch.setattr(module, "_collection_names",
                        lambda: {k: (v, "", "") for k, v in names.items()})
    monkeypatch.setattr(module, "_write_collection_name",
                        lambda cid, name: names.__setitem__(cid, name))
    monkeypatch.setattr(module.storage_mod, "get_storage", lambda: Store())

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("Shipments")
    sheet["A1"] = "Shipments"
    for index, header in enumerate(["Ref", "City", "Weight"]):
        sheet.cell(row=3, column=1 + index, value=header)
    for row in range(3):
        sheet.cell(row=4 + row, column=1, value=f"S-{row}")
        sheet.cell(row=4 + row, column=2, value="Newark")
        sheet.cell(row=4 + row, column=3, value=10 + row)
    path = tmp_path / "shipments.xlsx"
    workbook.save(path)
    payload = [("DA EXPRESS/shipments.xlsx", path.read_bytes())]

    first = module.ingest_files(payload, owner="me")
    second = module.ingest_files(payload, owner="me")

    assert first.collection_id != second.collection_id
    collections = {c.collection_id: c.name for c in module.list_collections()}
    assert collections[first.collection_id] == "DA EXPRESS"
    assert collections[second.collection_id] == "DA EXPRESS (2)"


def test_adding_to_a_folder_does_not_create_a_second_one(tmp_path, monkeypatch):
    """"Add files" must grow the folder, not clone it."""
    import openpyxl

    from app.engine import excel_sources as module

    saved: dict[str, Dataset] = {}
    names: dict[str, str] = {}
    store_dir = tmp_path / "store"
    store_dir.mkdir()

    class Store:
        def put(self, key, data):
            (store_dir / key.replace("/", "__")).write_bytes(data)

        def get(self, key):
            return (store_dir / key.replace("/", "__")).read_bytes()

    monkeypatch.setattr(module, "_ensure_tables", lambda: True)
    monkeypatch.setattr(module, "_save_dataset", lambda d: saved.__setitem__(d.dataset_id, d))
    monkeypatch.setattr(module, "list_datasets", lambda: list(saved.values()))
    monkeypatch.setattr(module, "get_dataset", lambda i: saved.get(i))
    monkeypatch.setattr(module, "load_definition", lambda fp: None)
    monkeypatch.setattr(module, "_save_definition", lambda *a, **k: None)
    monkeypatch.setattr(module, "known_concepts", lambda: [])
    monkeypatch.setattr(module, "_collection_names",
                        lambda: {k: (v, "", "") for k, v in names.items()})
    monkeypatch.setattr(module, "_write_collection_name",
                        lambda cid, name: names.__setitem__(cid, name))
    monkeypatch.setattr(module.storage_mod, "get_storage", lambda: Store())

    def book(name, headers, rows):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("S")
        ws["A1"] = name
        for i, h in enumerate(headers):
            ws.cell(row=3, column=1 + i, value=h)
        for r in range(rows):
            for c, h in enumerate(headers):
                ws.cell(row=4 + r, column=1 + c, value=f"{h}{r}")
        path = tmp_path / f"{name}.xlsx"
        wb.save(path)
        return path.read_bytes()

    same = ("DA EXPRESS/a.xlsx", book("a", ["Ref", "City", "Weight"], 3))
    first = module.ingest_files([same], owner="me")
    folder = first.collection_id

    # One file of the SAME shape, one of a new shape, added to that folder.
    added = module.ingest_files(
        [
            ("whatever/b.xlsx", book("b", ["Ref", "City", "Weight"], 2)),
            ("whatever/c.xlsx", book("c", ["Loan", "Rate"], 2)),
        ],
        owner="me",
        add_to_collection=folder,
    )

    assert added.collection_id == folder
    collections = [c for c in module.list_collections()]
    assert len(collections) == 1, "adding files must not create a second folder"
    grown = collections[0]
    # The matching shape joined the dataset already there; the new shape became
    # a second dataset inside the same folder.
    assert len(grown.datasets) == 2
    assert grown.file_count == 3
