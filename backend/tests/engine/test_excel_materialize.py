"""Materializer tests — a profiled workbook becoming queryable tables.

The cases here are the ones that decide whether a number a user later sees is
right: that same-shaped sheets stack into ONE table rather than twenty, that
subtotals never load beside the rows they sum, that text-encoded dates and
currency become real types, and that every row keeps a path back to the cell it
came from.
"""

from __future__ import annotations

import duckdb
import openpyxl
import pytest

from app.engine.excel_materialize import Batch, coerce, materialize, to_identifier
from app.engine.excel_profiler import profile_workbook


def _rent_roll_sheet(ws, name: str, rows: list[tuple]) -> None:
    ws["A1"] = name
    headers = ["Unit", "Tenant", "SF", "Rent/Mo", "Lease End"]
    for offset, text in enumerate(headers):
        ws.cell(row=3, column=1 + offset, value=text)
    for row_offset, values in enumerate(rows):
        for col_offset, value in enumerate(values):
            ws.cell(row=4 + row_offset, column=1 + col_offset, value=value)
        ws.cell(row=4 + row_offset, column=4).number_format = "$#,##0"


@pytest.fixture
def portfolio_workbook(tmp_path):
    """Three properties with identical rent-roll layouts."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _rent_roll_sheet(
        wb.create_sheet("01 Alpha Apartments"),
        "01 Alpha Apartments",
        [("101", "Ana Diaz", 900, 2400, "12/31/2026"),
         ("102", "Ben Cole", 850, 2250, "06/30/2027")],
    )
    _rent_roll_sheet(
        wb.create_sheet("02 Beta Apartments"),
        "02 Beta Apartments",
        [("201", "Cara Fox", 1100, 3100, "03/31/2027")],
    )
    _rent_roll_sheet(
        wb.create_sheet("03 Gamma Apartments"),
        "03 Gamma Apartments",
        [("301", "Dev Rao", 950, 2600, "09/30/2026")],
    )
    path = tmp_path / "portfolio.xlsx"
    wb.save(path)
    return path


def _load(path, tmp_path, mapping=None):
    database = tmp_path / "out.duckdb"
    result = materialize(path, profile_workbook(path), mapping, database)
    return result, duckdb.connect(str(database), read_only=True)


# ── identifiers and casting ───────────────────────────────────────────────────


@pytest.mark.parametrize(
    "name,expected",
    [
        ("Rent/Mo", "rent_mo"),
        ("Comments (Add notes as needed)", "comments_add_notes_as_needed"),
        ("# of Bed Rooms", "of_bed_rooms"),
        ("2026 Actual", "c_2026_actual"),
        ("", "col"),
    ],
)
def test_identifiers_are_safe_and_unquoted(name, expected):
    """A model writing SELECT \"Rent/Mo\" gets it wrong often enough to matter,
    so column names are normalised rather than preserved verbatim."""
    assert to_identifier(name) == expected


@pytest.mark.parametrize(
    "value,sql_type,expected",
    [
        ("$113,990", "DOUBLE", 113990.0),
        ("(1,234)", "DOUBLE", -1234.0),
        ("94%", "DOUBLE", 94.0),
        ("12/30/2023", "DATE", __import__("datetime").date(2023, 12, 30)),
        ("2023-12-30", "DATE", __import__("datetime").date(2023, 12, 30)),
        ("Yes", "BOOLEAN", True),
        ("No", "BOOLEAN", False),
        ("not a number", "DOUBLE", None),
        ("", "DOUBLE", None),
    ],
)
def test_text_encoded_values_become_real_types(value, sql_type, expected):
    assert coerce(value, sql_type) == expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("2.9x", 2.9),      # a coverage multiple, as written in a rent roll
        ("1,200 sqft", 1200.0),
        ("250 bps", 250.0),
        ("12 mo", 12.0),
        ("N/A", None),      # genuinely absent — must NOT become a number
        ("TBD", None),
        ("Verified", None),
    ],
)
def test_a_unit_written_after_the_number_does_not_lose_the_number(value, expected):
    """"2.9x" in an income-to-rent column parsed as nothing, so the whole column
    loaded NULL and the answer came back "no data available for this property".
    The unit is stripped only when what remains is a number, so "N/A" still
    fails as it should."""
    assert coerce(value, "DOUBLE") == expected


def test_a_column_that_never_parses_is_reported_rather_than_left_blank(tmp_path):
    """The failure this guards is silent: the column types cleanly, the rows
    load, and every figure is NULL with no reason given anywhere."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for offset, text in enumerate(["Unit", "Tenant", "Score"]):
        ws.cell(row=1, column=1 + offset, value=text)
    for index, label in enumerate(["Not scored", "Pending", "Unavailable", "None"]):
        ws.cell(row=2 + index, column=1, value=f"10{index}")
        ws.cell(row=2 + index, column=2, value="Someone")
        ws.cell(row=2 + index, column=3, value=label)
    path = tmp_path / "unparseable.xlsx"
    wb.save(path)

    mapping = {
        "regions": [
            {
                "suggested_table": "scores",
                "region_rows": "1-5",
                "sheets": ["Sheet"],
                "fields": [
                    {"source": "Score", "concept": "score",
                     "data_type": "number", "role": "measure"},
                ],
            }
        ]
    }
    result, conn = _load(path, tmp_path, mapping)
    try:
        table = max(result.tables, key=lambda t: t.row_count)
        assert "score" in table.unreadable_columns
    finally:
        conn.close()


def test_a_few_missing_values_are_not_reported_as_unreadable(tmp_path):
    """A rent roll legitimately holds "N/A" for a vacant unit. Flagging that
    would bury the real signal in noise."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for offset, text in enumerate(["Unit", "Tenant", "Ratio"]):
        ws.cell(row=1, column=1 + offset, value=text)
    for index, ratio in enumerate(["2.9x", "3.0x", "N/A", "3.4x", "2.8x"]):
        ws.cell(row=2 + index, column=1, value=f"10{index}")
        ws.cell(row=2 + index, column=2, value="Someone")
        ws.cell(row=2 + index, column=3, value=ratio)
    path = tmp_path / "mostly_fine.xlsx"
    wb.save(path)

    result, conn = _load(path, tmp_path)
    try:
        table = max(result.tables, key=lambda t: t.row_count)
        assert table.unreadable_columns == []
    finally:
        conn.close()


def test_an_unconvertible_cell_becomes_null_rather_than_failing():
    """One malformed cell must not fail a 2,000-row upload."""
    assert coerce("n/a", "DOUBLE") is None
    assert coerce("31/31/2023", "DATE") is None


# ── stacking ──────────────────────────────────────────────────────────────────


def test_same_shaped_sheets_stack_into_one_table(portfolio_workbook, tmp_path):
    """Three properties, one rent roll — otherwise every portfolio question
    needs an N-way UNION the model has to get right each time."""
    result, conn = _load(portfolio_workbook, tmp_path)
    try:
        tables = [t for t in result.tables if t.row_count == 4]
        assert len(tables) == 1
        assert len(tables[0].sheets) == 3
    finally:
        conn.close()


def test_every_row_records_where_it_came_from(portfolio_workbook, tmp_path):
    result, conn = _load(portfolio_workbook, tmp_path)
    try:
        table = max(result.tables, key=lambda t: t.row_count).name
        rows = conn.execute(
            f'SELECT source_sheet, source_row FROM "{table}" ORDER BY 1, 2'
        ).fetchall()
        assert rows[0] == ("01 Alpha Apartments", 4)
        assert {r[0] for r in rows} == {
            "01 Alpha Apartments",
            "02 Beta Apartments",
            "03 Gamma Apartments",
        }
    finally:
        conn.close()


def test_a_portfolio_question_is_one_group_by(portfolio_workbook, tmp_path):
    result, conn = _load(portfolio_workbook, tmp_path)
    try:
        table = max(result.tables, key=lambda t: t.row_count).name
        rows = conn.execute(
            f'SELECT source_sheet, SUM(rent_mo) FROM "{table}" GROUP BY 1 ORDER BY 2 DESC'
        ).fetchall()
        assert rows[0] == ("01 Alpha Apartments", 4650.0)
        assert len(rows) == 3
    finally:
        conn.close()


def test_dates_stored_as_text_are_queryable_as_dates(portfolio_workbook, tmp_path):
    result, conn = _load(portfolio_workbook, tmp_path)
    try:
        table = max(result.tables, key=lambda t: t.row_count).name
        expiring = conn.execute(
            f"SELECT COUNT(*) FROM \"{table}\" "
            "WHERE lease_end BETWEEN '2026-01-01' AND '2026-12-31'"
        ).fetchone()[0]
        assert expiring == 2
    finally:
        conn.close()


# ── subtotals ─────────────────────────────────────────────────────────────────


def test_a_total_row_is_never_loaded_beside_the_rows_it_sums(tmp_path):
    """Loading the total alongside its inputs doubles every figure derived
    from the table."""
    # Three columns, because two columns of label-and-value are structurally
    # indistinguishable from a form and get read as one.
    wb = openpyxl.Workbook()
    ws = wb.active
    for offset, text in enumerate(["Unit", "Tenant", "Rent/Mo"]):
        ws.cell(row=1, column=1 + offset, value=text)
    ws["A2"], ws["B2"], ws["C2"] = "101", "Ana Diaz", 1000
    ws["A3"], ws["B3"], ws["C3"] = "102", "Ben Cole", 2000
    ws["A4"], ws["C4"] = "Total", 3000
    path = tmp_path / "with_total.xlsx"
    wb.save(path)

    result, conn = _load(path, tmp_path)
    try:
        table = max(result.tables, key=lambda t: t.row_count).name
        total = conn.execute(f'SELECT SUM(rent_mo) FROM "{table}"').fetchone()[0]
        assert total == 3000.0  # not 6000
        assert conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] == 2
    finally:
        conn.close()


def test_derived_statement_lines_are_kept_out_of_the_input_table(tmp_path):
    """An operating statement's EGI and NOI are formulas over the lines above
    them, so they load into a separate table rather than beside their inputs."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01 Property"
    ws["A1"] = "INCOME"
    ws["A2"], ws["B2"] = "Gross potential rent", 1000
    ws["A3"], ws["B3"] = "Other income", 200
    ws["A4"], ws["B4"] = "Effective gross income", "=SUM(B2:B3)"
    path = tmp_path / "statement.xlsx"
    wb.save(path)

    result, conn = _load(path, tmp_path)
    try:
        names = {t.name for t in result.tables}
        table = next(iter(names))
        columns = {
            row[0]
            for row in conn.execute(
                "SELECT column_name FROM information_schema.columns "
                f"WHERE table_name = '{table}'"
            ).fetchall()
        }
        assert "gross_potential_rent" in columns
        assert "effective_gross_income" not in columns
        derived = conn.execute(
            "SELECT column_name FROM information_schema.columns "
            f"WHERE table_name = '{table}_derived'"
        ).fetchall()
        assert any("effective" in row[0] for row in derived)
    finally:
        conn.close()


# ── many files, one dataset ───────────────────────────────────────────────────


def _monthly_rent_roll(tmp_path, month: int, rent: int, name: str | None = None):
    """One month of an otherwise identical rent roll."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = f"Rent Roll — as of 2026-{month:02d}-01"
    for offset, text in enumerate(["Unit", "Tenant", "Rent"]):
        ws.cell(row=3, column=1 + offset, value=text)
    for index in range(3):
        ws.cell(row=4 + index, column=1, value=f"10{index}")
        ws.cell(row=4 + index, column=2, value="Someone")
        ws.cell(row=4 + index, column=3, value=rent)
    path = tmp_path / (name or f"RentRoll_2026-{month:02d}.xlsx")
    wb.save(path)
    return path


def _load_batch(paths_and_rents, tmp_path):
    """Load a series of files into ONE dataset, as a folder upload would."""
    database = tmp_path / "dataset.duckdb"
    for index, path in enumerate(paths_and_rents):
        profile = profile_workbook(path)
        materialize(
            path,
            profile,
            None,
            database,
            batch=Batch(
                source_file=path.name, as_of=profile.as_of, append=index > 0
            ),
        )
    return duckdb.connect(str(database), read_only=True)


def test_files_of_the_same_shape_stack_into_one_trendable_table(tmp_path):
    """The whole point of the dataset model: twelve monthly files become one
    table you can GROUP BY, not twelve isolated uploads."""
    paths = [_monthly_rent_roll(tmp_path, m, 2000 + m * 25) for m in range(1, 13)]
    conn = _load_batch(paths, tmp_path)
    try:
        table = conn.execute(
            "SELECT table_name FROM information_schema.tables"
        ).fetchone()[0]
        assert conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] == 36
        months = conn.execute(
            f'SELECT COUNT(DISTINCT as_of) FROM "{table}"'
        ).fetchone()[0]
        assert months == 12

        trend = conn.execute(
            f'SELECT as_of, AVG(rent) FROM "{table}" GROUP BY 1 ORDER BY 1'
        ).fetchall()
        assert trend[0][1] < trend[-1][1], "rent should rise across the year"
    finally:
        conn.close()


def test_re_uploading_a_file_replaces_its_rows_rather_than_doubling_them(tmp_path):
    """Someone will upload a corrected March. Without replace-by-file that
    month silently doubles and every total for it is wrong."""
    paths = [_monthly_rent_roll(tmp_path, m, 2000) for m in (1, 2, 3)]
    conn = _load_batch(paths, tmp_path)
    table = conn.execute(
        "SELECT table_name FROM information_schema.tables"
    ).fetchone()[0]
    conn.close()

    corrected = _monthly_rent_roll(tmp_path, 3, 9999)
    profile = profile_workbook(corrected)
    materialize(
        corrected,
        profile,
        None,
        tmp_path / "dataset.duckdb",
        batch=Batch(source_file=corrected.name, as_of=profile.as_of, append=True),
    )

    conn = duckdb.connect(str(tmp_path / "dataset.duckdb"), read_only=True)
    try:
        assert conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] == 9
        march = conn.execute(
            f"SELECT DISTINCT rent FROM \"{table}\" WHERE as_of = '2026-03-01'"
        ).fetchall()
        assert march == [(9999.0,)]
    finally:
        conn.close()


def test_every_row_records_which_file_it_came_from(tmp_path):
    """Citation across files, not just across sheets."""
    paths = [_monthly_rent_roll(tmp_path, m, 2000) for m in (1, 2)]
    conn = _load_batch(paths, tmp_path)
    try:
        table = conn.execute(
            "SELECT table_name FROM information_schema.tables"
        ).fetchone()[0]
        files = {
            row[0]
            for row in conn.execute(
                f'SELECT DISTINCT source_file FROM "{table}"'
            ).fetchall()
        }
        assert files == {"RentRoll_2026-01.xlsx", "RentRoll_2026-02.xlsx"}
    finally:
        conn.close()


def test_a_file_with_no_period_still_loads(tmp_path):
    """An undated file answers everything except a trend. Refusing it would be
    worse than carrying it with a null period."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for offset, text in enumerate(["Unit", "Tenant", "Rent"]):
        ws.cell(row=1, column=1 + offset, value=text)
    ws["A2"], ws["B2"], ws["C2"] = "101", "Someone", 2000
    path = tmp_path / "Rent Roll Final.xlsx"
    wb.save(path)

    profile = profile_workbook(path)
    assert profile.as_of is None

    database = tmp_path / "undated.duckdb"
    result = materialize(
        path, profile, None, database, batch=Batch(source_file=path.name)
    )
    assert result.total_rows == 1


# ── mapping ───────────────────────────────────────────────────────────────────


def test_the_mapping_overrides_a_numeric_looking_identifier(
    portfolio_workbook, tmp_path
):
    """Unit numbers look numeric and must not become a summable column. The
    profiler cannot know that; the mapper can, and it wins."""
    mapping = {
        "regions": [
            {
                "suggested_table": "unit_rent_roll",
                "region_rows": "3-5",
                "sheets": ["01 Alpha Apartments"],
                "grain": "One unit",
                "purpose": "Rent roll",
                "fields": [
                    {
                        "source": "Unit",
                        "concept": "unit_number",
                        "data_type": "text",
                        "role": "identifier",
                    },
                    {
                        "source": "Rent/Mo",
                        "concept": "monthly_rent",
                        "data_type": "currency",
                        "role": "measure",
                    },
                ],
            }
        ]
    }
    result, conn = _load(portfolio_workbook, tmp_path, mapping)
    try:
        table = max(result.tables, key=lambda t: t.row_count)
        assert table.name == "unit_rent_roll"
        types = dict(
            conn.execute(
                "SELECT column_name, data_type FROM information_schema.columns "
                f"WHERE table_name = '{table.name}'"
            ).fetchall()
        )
        assert types["unit_number"] == "VARCHAR"
        assert types["monthly_rent"] == "DOUBLE"
    finally:
        conn.close()


def test_a_workbook_still_loads_without_any_mapping(portfolio_workbook, tmp_path):
    """A model outage should degrade the names, not lose the data."""
    result, conn = _load(portfolio_workbook, tmp_path, None)
    try:
        assert result.total_rows == 4
        assert not result.warnings
    finally:
        conn.close()


# -- files of one shape share one table ---------------------------------------


def _deal_sheet(ws, title: str, prefix: str, count: int) -> None:
    """A deal register: same 6 columns, different sheet name and length."""
    ws["A1"] = f"{title} - Deal Pipeline"
    headers = ["Deal ID", "City", "Property Type", "Deal Stage",
               "Close Price", "Commission Amount"]
    for offset, text in enumerate(headers):
        ws.cell(row=3, column=1 + offset, value=text)
    for index in range(count):
        for col_offset, value in enumerate(
            (f"{prefix}-{1001 + index}", "Newark", title, "Closed",
             100000 + index, 5000 + index)
        ):
            ws.cell(row=4 + index, column=1 + col_offset, value=value)


def _deal_workbook(tmp_path, filename: str, title: str, prefix: str, count: int):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _deal_sheet(wb.create_sheet(title), title, prefix, count)
    path = tmp_path / filename
    wb.save(path)
    return path


@pytest.fixture
def deal_files(tmp_path):
    """Three same-shaped registers that agree on nothing but their headers.

    Different sheet names and different row counts are what separate one file of
    a dataset from the next.
    """
    return [
        _deal_workbook(tmp_path, "commercial.xlsx", "Commercial", "C", 3),
        _deal_workbook(tmp_path, "industrial.xlsx", "Industrial", "I", 5),
        _deal_workbook(tmp_path, "land.xlsx", "Land", "L", 2),
    ]


def _stack(paths, tmp_path, mapping=None):
    database = tmp_path / "stacked.duckdb"
    results = [
        materialize(
            path, profile_workbook(path), mapping, database,
            batch=Batch(source_file=path.name, append=index > 0),
        )
        for index, path in enumerate(paths)
    ]
    return results, duckdb.connect(str(database), read_only=True)


def test_same_shaped_files_stack_into_one_table(deal_files, tmp_path):
    """Files of one shape share a table, whatever their sheets are called.

    A table per file leaves the dataset describing only the last one, and every
    question is then answered from a fraction of the rows.
    """
    results, conn = _stack(deal_files, tmp_path)
    try:
        names = {t.name for result in results for t in result.tables}
        assert len(names) == 1, f"one shape must be one table, got {names}"
        table = names.pop()
        assert conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] == 10
        assert conn.execute(
            f'SELECT COUNT(DISTINCT source_file) FROM "{table}"'
        ).fetchone()[0] == 3
    finally:
        conn.close()


def test_the_table_name_never_comes_from_the_sheet_name(deal_files, tmp_path):
    """A name built from the sheet or the row range cannot survive the next file."""
    results, conn = _stack(deal_files, tmp_path)
    try:
        names = [t.name for result in results for t in result.tables]
        for name in names:
            for sheet_name in ("commercial", "industrial", "land"):
                assert sheet_name not in name, f"{name} was named after a sheet"
        # The three files span rows 3-6, 3-8 and 3-5. A name carrying the range
        # would differ between them; one built from the columns cannot.
        assert len(set(names)) == 1
    finally:
        conn.close()


def test_a_mapping_made_from_one_file_names_the_tables_of_all_of_them(
    deal_files, tmp_path
):
    """The mapper only ever sees the first file of a shape.

    Its record therefore names one sheet and one row range, so matching on
    either sends every later file elsewhere. Matching on the column names is
    what makes a saved definition reusable.
    """
    mapping = {
        "regions": [
            {
                "suggested_table": "deal_pipeline",
                "region_rows": "3-6",
                "sheets": ["Commercial"],
                "grain": "One deal",
                "purpose": "Deals in the pipeline",
                "fields": [
                    {"source": "Deal ID", "concept": "deal_id", "data_type": "text"},
                    {"source": "Close Price", "concept": "close_price",
                     "data_type": "currency"},
                ],
            }
        ]
    }
    results, conn = _stack(deal_files, tmp_path, mapping)
    try:
        assert {t.name for r in results for t in r.tables} == {"deal_pipeline"}
        total = conn.execute(
            "SELECT COUNT(*), SUM(close_price) FROM deal_pipeline"
        ).fetchone()
        assert total[0] == 10
    finally:
        conn.close()


def test_a_summary_of_only_totals_is_reported_as_a_real_table(tmp_path):
    """A block of only totals is one table, not an empty shell and a hidden twin.

    Split, its figures are in the database and named nowhere, so nothing can
    select them and the library shows a table with no rows.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    _deal_sheet(ws, "Commercial", "C", 3)
    ws["A9"] = "Summary"
    for offset, (label, value) in enumerate(
        (("Total Deals", 3), ("Total Closed Volume", 300003))
    ):
        ws.cell(row=10 + offset, column=1, value=label)
        ws.cell(row=10 + offset, column=2, value=value)
    path = tmp_path / "with_summary.xlsx"
    wb.save(path)

    result, conn = _load(path, tmp_path)
    try:
        assert all(t.row_count > 0 for t in result.tables), (
            f"no empty shells: {[(t.name, t.row_count) for t in result.tables]}"
        )
        summary = next(t for t in result.tables if "total_deals" in t.name)
        assert "total_closed_volume" in summary.columns
    finally:
        conn.close()
