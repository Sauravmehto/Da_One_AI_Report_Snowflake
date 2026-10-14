"""Documents becoming a searchable, citable graph.

The cases that decide whether an answer drawn from a document can be trusted:
that decoration never reaches the text, that a scan is reported rather than
indexed as an empty document, that one property resolves to one entity however
its name is written, and that every passage leads back to a file and a page.
"""

from __future__ import annotations

import duckdb
import pytest

from app.engine import document_store as store
from app.engine.document_parser import parse_document, parse_pdf


@pytest.fixture
def conn():
    connection = duckdb.connect(":memory:")
    store.ensure_schema(connection)
    yield connection
    connection.close()


def _pdf(tmp_path, lines, *, watermark=None, heading=None, name="doc.pdf"):
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    if heading:
        page.insert_text((60, 50), heading, fontsize=24)
    if watermark:
        page.insert_text((300, 400), watermark, fontsize=30, rotate=90)
    y = 100
    for line in lines:
        page.insert_text((60, y), line, fontsize=10)
        y += 16
    path = tmp_path / name
    document.save(path)
    document.close()
    return path


# -- reading a page -----------------------------------------------------------


def test_a_rotated_watermark_never_reaches_the_text(tmp_path):
    """A watermark otherwise bleeds characters into the values beside it."""
    path = _pdf(
        tmp_path,
        ["Loan Amount $12,400,000", "Interest Rate 5.85%"],
        watermark="CONFIDENTIAL DRAFT",
    )
    parsed = parse_pdf(path)
    assert "CONFIDENTIAL" not in parsed.text
    assert "12,400,000" in parsed.text


def test_an_oversized_heading_is_decoration_not_body(tmp_path):
    """Decoration is judged against the page's own body size.

    The page needs real body text for that to mean anything: "biggest size on
    the page" is not the test, "far bigger than what most of this page is set
    in" is.
    """
    body = [f"Clause {i}. The parties agree to the terms set out herein." for i in range(12)]
    path = _pdf(tmp_path, body, heading="ANNUAL REPORT 2026")
    parsed = parse_pdf(path)
    assert "ANNUAL REPORT" not in parsed.text
    assert "Clause 0." in parsed.text
    assert "Clause 11." in parsed.text


def test_a_page_with_no_text_layer_is_reported_as_scanned(tmp_path):
    """A scan extracts to nothing and would otherwise look like an empty file."""
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 300, 300))
    path = tmp_path / "scan.pdf"
    document.save(path)
    document.close()

    parsed = parse_pdf(path)
    assert parsed.is_scanned
    assert parsed.scanned_pages == [1]


def test_a_readable_pdf_is_not_called_scanned(tmp_path):
    path = _pdf(tmp_path, ["A" * 80, "B" * 80])
    assert not parse_pdf(path).is_scanned


# -- other formats ------------------------------------------------------------


def test_a_word_document_is_read_into_pages(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Termination. Ninety days written notice is required.")
    document.add_paragraph("Governing law is New York.")
    path = tmp_path / "agreement.docx"
    document.save(path)

    parsed = parse_document(path)
    assert parsed.file_format == "docx"
    assert "Ninety days" in parsed.text
    assert "New York" in parsed.text


def test_plain_text_is_read(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("First para.\n\nSecond para.", encoding="utf-8")
    parsed = parse_document(path)
    assert parsed.file_format == "txt"
    assert "Second para." in parsed.text


def test_an_unsupported_format_says_so(tmp_path):
    path = tmp_path / "thing.pptx"
    path.write_bytes(b"x")
    with pytest.raises(ValueError, match="Cannot read"):
        parse_document(path)


# -- chunking -----------------------------------------------------------------


def test_a_short_page_is_one_passage():
    assert store.split_into_chunks("A short clause.") == ["A short clause."]


def test_a_long_page_is_split_with_overlap():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 60 for i in range(20))
    chunks = store.split_into_chunks(text)
    assert len(chunks) > 1
    assert all(len(c) <= store.CHUNK_CHARS + store.CHUNK_OVERLAP for c in chunks)
    # Nothing is lost at a boundary: every paragraph survives somewhere.
    joined = " ".join(chunks)
    for i in range(20):
        assert f"Paragraph {i}." in joined


def test_an_empty_page_produces_no_passages():
    assert store.split_into_chunks("   \n  ") == []


# -- entity resolution --------------------------------------------------------


@pytest.mark.parametrize(
    "left,right",
    [
        ("Cherry Creek Apartments", "Cherry Creek Apts LLC"),
        ("The Meridian Building", "Meridian"),
        ("Oak Street Properties, Inc.", "Oak Street"),
    ],
)
def test_one_property_resolves_to_one_entity(conn, left, right):
    """The same building is named differently in every file that mentions it."""
    assert store.resolve_entity(conn, left) == store.resolve_entity(conn, right)


def test_different_properties_stay_apart(conn):
    assert store.resolve_entity(conn, "Cherry Creek") != store.resolve_entity(
        conn, "Willow Bend"
    )


def test_an_alias_can_be_added_later(conn):
    entity = store.resolve_entity(conn, "Cherry Creek")
    store.add_alias(conn, entity, "CCSR Holdings")
    assert store.resolve_entity(conn, "CCSR Holdings") == entity


# -- writing and citing -------------------------------------------------------


def _write(conn, tmp_path, name="mortgage.pdf", lines=None):
    lines = lines or [
        "Mortgage Agreement between Cherry Creek Apts LLC and Meridian Capital.",
        "Termination requires ninety days written notice.",
    ]
    parsed = parse_pdf(_pdf(tmp_path, lines, name=name))
    return store.write_document(
        conn,
        parsed,
        doc_type="mortgage_agreement",
        facts=[
            store.Fact(
                "Cherry Creek Apartments", "loan_amount",
                value_num=12_400_000, unit="usd", source_page=1,
            )
        ],
        edges=[
            store.Edge(
                "Cherry Creek Apartments", "borrows_from", "Meridian Capital",
                source_page=1,
            )
        ],
    )


def test_a_document_lands_as_passages_facts_and_edges(conn, tmp_path):
    outcome = _write(conn, tmp_path)
    assert outcome.chunk_count >= 1
    assert outcome.fact_count == 1
    assert outcome.edge_count == 1
    assert not outcome.warnings


def test_every_passage_carries_its_file_and_page(conn, tmp_path):
    _write(conn, tmp_path)
    rows = conn.execute("SELECT source_file, source_page FROM chunks").fetchall()
    assert rows
    assert all(r[0] == "mortgage.pdf" and r[1] == 1 for r in rows)


def test_reloading_a_document_replaces_it(conn, tmp_path):
    """A corrected file must replace the old one, not double it."""
    _write(conn, tmp_path)
    before = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    _write(conn, tmp_path)
    after = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    assert after == before
    assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 1
    assert conn.execute("SELECT COUNT(*) FROM facts").fetchone()[0] == 1


def test_a_scanned_document_warns_rather_than_indexing_nothing(conn, tmp_path):
    import pymupdf

    document = pymupdf.open()
    document.new_page()
    path = tmp_path / "scan.pdf"
    document.save(path)
    document.close()

    outcome = store.write_document(conn, parse_pdf(path))
    assert outcome.is_scanned
    assert outcome.chunk_count == 0
    assert any("scan" in w for w in outcome.warnings)


# -- searching ----------------------------------------------------------------


def test_search_finds_a_clause_and_cites_it(conn, tmp_path):
    _write(conn, tmp_path)
    store.build_index(conn)
    hits = store.search(conn, "termination notice")
    assert hits
    assert hits[0].source_file == "mortgage.pdf"
    assert hits[0].source_page == 1


def test_search_can_be_scoped_to_one_file(conn, tmp_path):
    _write(conn, tmp_path, name="a.pdf")
    _write(conn, tmp_path, name="b.pdf")
    store.build_index(conn)
    hits = store.search(conn, "termination", source_file="b.pdf")
    assert hits
    assert {h.source_file for h in hits} == {"b.pdf"}


def test_search_still_works_without_the_index(conn, tmp_path):
    """The FTS extension downloads on first use, so it can be unavailable."""
    _write(conn, tmp_path)
    hits = store._fallback_search(conn, "termination notice", 5, None)
    assert hits
    assert hits[0].source_file == "mortgage.pdf"


def test_an_empty_query_finds_nothing(conn, tmp_path):
    _write(conn, tmp_path)
    assert store.search(conn, "   ") == []


def test_the_summary_describes_what_is_stored(conn, tmp_path):
    _write(conn, tmp_path)
    summary = store.document_summary(conn)
    assert len(summary) == 1
    assert summary[0]["source_file"] == "mortgage.pdf"
    assert summary[0]["doc_type"] == "mortgage_agreement"
    assert summary[0]["chunks"] >= 1


def test_a_database_with_no_documents_says_so(conn):
    assert not store.has_documents(conn)
