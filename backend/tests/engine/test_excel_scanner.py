"""Unit tests for excel_template_scanner — pure classification, no Snowflake."""

import openpyxl

from app.engine import excel_template_scanner as scanner


def _build_workbook(data_rows: int):
    """A template with a large 'data' sheet and a 'Summary' formula sheet that
    exercises all three formula tiers."""
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    data["A1"] = "Region"
    data["B1"] = "Amount"
    for i in range(data_rows):
        data.cell(row=2 + i, column=1).value = f"R{i}"
        data.cell(row=2 + i, column=2).value = 100 + i

    summary = wb.create_sheet("Summary")
    summary["A2"] = "Total"
    summary["B2"] = "=SUM(Data!B2:B100)"      # Tier 1: aggregate over a range
    summary["A3"] = "Snapshot"
    summary["B3"] = "=Data!B5"                 # Tier 2: single cross-sheet cell
    summary["A4"] = "Grossed up"
    summary["B4"] = "=B2*1.05"                 # Tier 3: static same-sheet arithmetic
    summary["A5"] = "Lookup"
    summary["B5"] = "=VLOOKUP(A5,Data!A:B,2,0)"  # Tier 2: unrecognized function
    return wb


def test_classify_formula_tiers():
    assert scanner.classify_formula("=SUM(Data!B2:B100)")[0] == "tier_1"
    assert scanner.classify_formula("=SUMIFS(A:A,B:B,1)")[0] == "tier_1"
    assert scanner.classify_formula("=Data!B5")[0] == "tier_2"
    assert scanner.classify_formula("=VLOOKUP(A5,Data!A:B,2,0)")[0] == "tier_2"
    assert scanner.classify_formula("=B2*1.05")[0] == "tier_3"
    assert scanner.classify_formula("=1+2")[0] == "tier_3"


def test_sheet_roles_and_tier_counts():
    scan = scanner.scan_workbook(_build_workbook(20), "demo")
    assert scan["tab_roles"]["Data"] == "data"
    assert scan["tab_roles"]["Summary"] == "formula"

    tiers = scan["formula_classification"]
    tier1 = {c["cell"] for c in tiers["tier_1"]}
    tier2 = {c["cell"] for c in tiers["tier_2"]}
    tier3 = {c["cell"] for c in tiers["tier_3"]}
    assert tier1 == {"B2"}
    assert tier2 == {"B3", "B5"}
    assert tier3 == {"B4"}
    # Tier-1 cells carry the range so a run can re-extend them.
    assert tiers["tier_1"][0]["ranges"] == ["B2:B100"]


def test_fingerprint_is_invariant_to_row_count():
    a = scanner.scan_workbook(_build_workbook(20), "demo")
    b = scanner.scan_workbook(_build_workbook(200), "demo")
    assert a["fingerprint"] == b["fingerprint"]


def test_fingerprint_changes_with_shape():
    base = scanner.scan_workbook(_build_workbook(20), "demo")
    wb = _build_workbook(20)
    wb["Summary"]["B6"] = "=SUM(Data!A2:A100)"   # a new formula changes the shape
    changed = scanner.scan_workbook(wb, "demo")
    assert base["fingerprint"] != changed["fingerprint"]


def test_small_zero_formula_data_table_is_data_not_parameter():
    """Regression: the minimal synthetic template from the base64-upload bug — a
    tiny, zero-formula record table (header + 3 rows, only 8 populated cells) must
    be 'data', NOT 'parameter'. Cell count alone can't tell them apart; column
    type-consistency does."""
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "DATA"
    data["A1"], data["B1"] = "Region", "Revenue"
    for i, (region, revenue) in enumerate([("North", 100), ("South", 200), ("East", 150)]):
        data.cell(row=2 + i, column=1).value = region
        data.cell(row=2 + i, column=2).value = revenue

    summary = wb.create_sheet("SUMMARY")
    summary["A1"], summary["B1"] = "Total Rev", "=SUM(DATA!B2:B4)"
    summary["A2"], summary["B2"] = "Region Count", "=COUNTA(DATA!A2:A4)"

    scan = scanner.scan_workbook(wb, "minimal")
    assert scan["tab_roles"]["DATA"] == "data"          # the bug: was "parameter"
    assert scan["tab_roles"]["SUMMARY"] == "formula"


def test_scattered_label_value_sheet_is_parameter_not_data():
    """The other side of the distinction: a label/value config sheet with a
    section header, a blank row, and a TYPE-MIXED value column is 'parameter',
    even though it also has zero formulas and a repeating 2-column shape."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASSUMPTIONS"
    ws["A1"] = "Fund Parameters"                        # section header (1 cell)
    ws["A3"], ws["B3"] = "Field", "Value"               # header row
    pairs = [("Fund Name", "Meridian I"), ("Vintage", 2021), ("Target IRR", 0.15),
             ("Currency", "USD"), ("Committed", 500000000)]   # value col mixes str/num
    for i, (label, value) in enumerate(pairs):
        ws.cell(row=4 + i, column=1).value = label
        ws.cell(row=4 + i, column=2).value = value

    scan = scanner.scan_workbook(wb, "assumptions")
    assert scan["tab_roles"]["ASSUMPTIONS"] == "parameter"
