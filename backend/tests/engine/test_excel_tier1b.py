"""Tier-1b row-mirror detection + STEP-2 value-sourcing tests.

Covers: generic (non-template-specific) detection on a synthetic workbook, the
true-positive/true-negative on the real Meridian template, and a guard that the
detection LOGIC hard-codes no template-specific names.

Local test fixture: the real-template test needs
`test_fixtures/excel/CreditFund_RiskDashboard (1).xlsx`. That workbook is real
portfolio data and is gitignored (NOT committed) — it must be present locally.
If it is missing the test FAILS LOUDLY (it does not silently skip) so it is
obvious the fixture is absent and why. See README "Local test fixtures"."""

import base64
import os
import re

import openpyxl
import pytest

from app.engine import excel_template_scanner as scanner
from app.engine import excel_reports as er


def _require_sample():
    """Fail loudly (not skip) if the gitignored local fixture is absent."""
    if not os.path.isfile(SAMPLE):
        pytest.fail(
            "Required local fixture is missing:\n  " + SAMPLE +
            "\nThis real portfolio-data workbook is gitignored (not committed) and "
            "must be present locally to run this test. Place it at that path. "
            "See README 'Local test fixtures'. (Failing loudly rather than "
            "silently skipping, by design.)"
        )

# backend/tests/engine/ -> backend/ (three levels up from this file)
SAMPLE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "test_fixtures", "excel", "CreditFund_RiskDashboard (1).xlsx")


# --- shift primitive ---------------------------------------------------------

def test_shift_formula_rows_relative_absolute_and_functions():
    assert scanner.shift_formula_rows("=LOAN!A3", 1) == "=LOAN!A4"
    assert scanner.shift_formula_rows("=G26-F26", 1) == "=G27-F27"
    assert scanner.shift_formula_rows('=IFERROR(LOAN!AH3,"")', 2) == '=IFERROR(LOAN!AH5,"")'
    # absolute row and bare numeric literals are untouched
    assert scanner.shift_formula_rows("=SUM($A$3:$A$7)*12", 1) == "=SUM($A$3:$A$7)*12"
    # a function name that looks like a cell ref (LOG10) must NOT shift
    assert scanner.shift_formula_rows("=LOG10(A2)", 1) == "=LOG10(A3)"


# --- generic synthetic workbook (STEP 3) -------------------------------------

def _synthetic():
    """A workbook with completely different names than the Meridian template:
    a same-sheet mirror block (SalesRaw!D), a cross-sheet mirror block
    (Summary!A), and a deliberately fragile scrambled-offset column (Summary!C)."""
    wb = openpyxl.Workbook()
    raw = wb.active
    raw.title = "SalesRaw"
    for col, name in zip("ABCDE", ["Order", "Qty", "Price", "Line Total", "Margin"]):
        raw[f"{col}1"] = name
    for i in range(8):                       # rows 2..9
        r = 2 + i
        raw[f"A{r}"] = f"ORD{i}"
        raw[f"B{r}"] = 10 + i
        raw[f"C{r}"] = 3 + i
        raw[f"D{r}"] = f"=B{r}*C{r}"         # same-sheet row-mirror block
        raw[f"E{r}"] = f"=D{r}-C{r}"

    summ = wb.create_sheet("Summary")
    summ["A1"] = "Mirror of orders"
    for i in range(8):                       # rows 3..10, cross-sheet mirror
        summ[f"A{3 + i}"] = f"=SalesRaw!A{2 + i}"
    # scrambled offsets 5,6,2,3,4 -> NOT a clean +1 shift (fragile, must stay T2)
    scrambled = [5, 6, 2, 3, 4]
    for i, off in enumerate(scrambled):      # rows 3..7
        summ[f"C{3 + i}"] = f"=SalesRaw!D{off}"
    return wb


def test_generic_row_mirror_detection():
    scan = scanner.scan_workbook(_synthetic(), "synthetic")
    blocks = {(b["sheet"], b["column"]): b for b in scan["formula_classification"]["tier_1b"]}

    assert ("SalesRaw", "D") in blocks
    assert blocks[("SalesRaw", "D")]["rows"] == 8
    assert ("SalesRaw", "E") in blocks
    assert ("Summary", "A") in blocks
    assert blocks[("Summary", "A")]["start_row"] == 3 and blocks[("Summary", "A")]["end_row"] == 10

    # the scrambled column is NOT a block, and its cells are flagged Tier 2
    assert ("Summary", "C") not in blocks
    t2_cells = {c["cell"] for c in scan["formula_classification"]["tier_2"]}
    assert {"C3", "C4", "C5"}.issubset(t2_cells)


def test_totals_row_under_a_block_does_not_kill_the_block():
    """Regression: a totals row directly beneath a mirror block is a natural
    BOUNDARY (different formula shape), not a violation. Previously it made the
    whole run inconsistent and the block was silently missed — which meant those
    cells never regenerated and the report went stale."""
    wb = openpyxl.Workbook()
    raw = wb.active
    raw.title = "Src"
    for i in range(5):
        raw.cell(row=2 + i, column=1).value = 10 + i

    rep = wb.create_sheet("Rep")
    for i in range(5):                       # rows 2..6 mirror Src!A2..A6
        rep[f"B{2 + i}"] = f"=Src!A{2 + i}"
    rep["B7"] = "=SUM(B2:B6)"                # totals row, contiguous with the block

    scan = scanner.scan_workbook(wb, "totals")
    blocks = {(b["sheet"], b["column"]): b for b in scan["formula_classification"]["tier_1b"]}
    assert ("Rep", "B") in blocks
    assert blocks[("Rep", "B")]["start_row"] == 2 and blocks[("Rep", "B")]["end_row"] == 6
    # the totals row itself is classified on its own, as a Tier-1 aggregate
    assert any(c["cell"] == "B7" for c in scan["formula_classification"]["tier_1"])


def test_tier2_risk_split_high_only_for_pinned_records():
    # pinned to one record -> high risk; aggregates/same-sheet -> low
    assert scanner.fixed_record_refs("=IFERROR(D!L6/SUM(D!L3:L7),0)") == ["D!L6"]
    assert scanner.fixed_record_refs("=IFERROR(AVERAGEIF(D!E3:E7,A17,D!S3:S7),\"-\")") == []
    assert scanner.fixed_record_refs('=IF(D17/SUM(D!L3:L7)*100>20,"BREACH","OK")') == []


def test_detection_logic_has_no_template_specific_names():
    """The genericness guarantee: no Meridian/template identifiers in the actual
    detection code (fixtures may use them, source logic may not)."""
    banned = ["RISK_DASHBOARD", "LOAN_DATA", "EXPOSURE_ANALYSIS", "TIMESERIES",
              "ASSUMPTIONS", "MERIDIAN", "DAYS_TO_MATURITY", "ALL_IN_YIELD",
              "ANNUAL_INCOME"]
    for mod in (scanner, er):
        src = open(mod.__file__, encoding="utf-8").read().upper()
        hits = [b for b in banned if b in src]
        assert not hits, f"{mod.__name__} references template-specific names: {hits}"


# --- STEP 2 value-sourcing name match ---------------------------------------

def test_value_source_name_match_is_fuzzy():
    header_by_col = {"AH": "Days to Maturity", "AI": "All-In Yield",
                     "AJ": "Annual Income ($)", "B": "Units"}
    q = ["SECURITY_ID", "DAYS_TO_MATURITY", "ALL_IN_YIELD", "ANNUAL_INCOME"]
    matched = er.resolve_value_sourced_columns(header_by_col, q)
    assert matched == {"AH": "DAYS_TO_MATURITY", "AI": "ALL_IN_YIELD", "AJ": "ANNUAL_INCOME"}


def test_regenerate_block_formulas_fill_shifts():
    assert er.regenerate_block_formulas("=SalesRaw!A2", 3) == \
        ["=SalesRaw!A2", "=SalesRaw!A3", "=SalesRaw!A4"]


def test_regenerate_tier1_only_extends_data_sheet_ranges():
    # a range qualified with the data sheet extends; a same-sheet range on a
    # dashboard sheet (a fixed sector table) is left alone.
    f = "=SUM(LOAN!K3:K7)"
    assert er.regenerate_tier1_formula(f, 12, "LOAN", on_data_sheet=False) == "=SUM(LOAN!K3:K12)"
    g = "=SUM(B17:B21)"
    assert er.regenerate_tier1_formula(g, 12, "LOAN", on_data_sheet=False) == g


# --- real Meridian template: true positive + true negative -------------------

def test_real_template_true_positive_and_negative():
    _require_sample()
    scan = scanner.scan_file(SAMPLE)

    # tab roles: record tables are "data", the assumptions sheet is "parameter",
    # the dashboards are "formula" (guards the data-vs-parameter fix).
    roles = scan["tab_roles"]
    assert roles["LOAN_DATA"] == "data"
    assert roles["TIMESERIES"] == "data"
    assert roles["ASSUMPTIONS"] == "parameter"
    assert roles["RISK_DASHBOARD"] == "formula"

    blocks = scan["formula_classification"]["tier_1b"]
    block_cols = {(b["sheet"], b["column"]) for b in blocks}

    # TRUE POSITIVE: the RISK_DASHBOARD row-mirror block (col A, rows 26-30).
    assert ("RISK_DASHBOARD", "A") in block_cols
    a_block = next(b for b in blocks if b["sheet"] == "RISK_DASHBOARD" and b["column"] == "A")
    assert a_block["start_row"] == 26 and a_block["end_row"] == 30
    # the same-sheet arithmetic (H) and function-wrapped (L) mirrors too.
    assert ("RISK_DASHBOARD", "H") in block_cols
    assert ("RISK_DASHBOARD", "L") in block_cols
    # LOAN_DATA derived columns are row-mirror blocks.
    assert ("LOAN_DATA", "AH") in block_cols

    # TRUE NEGATIVE: the fragile scrambled single-cell column E17:E21 is NOT a
    # block, and E17 is flagged Tier 2, never silently auto-regenerated.
    all_block_cells = {c for b in blocks for c in b["cells"]}
    assert "E17" not in all_block_cells
    t2_cells = {c["cell"] for c in scan["formula_classification"]["tier_2"]}
    assert "E17" in t2_cells


def test_scan_via_base64_matches_file_path(monkeypatch, tmp_path):
    """The two input modes must produce IDENTICAL classification — same tier
    counts, same fingerprint, same tab roles, same per-cell classification —
    proving base64 upload and local path are equivalent, not merely both 'work'."""
    _require_sample()
    # isolate from Snowflake and real storage: no fingerprint lookup, temp uploads
    monkeypatch.setattr(er, "_find_template_by_fingerprint", lambda fp: None)
    monkeypatch.setattr(er, "EXCEL_STORAGE_DIR", tmp_path)

    via_path = er.scan_template(file_path=SAMPLE)

    with open(SAMPLE, "rb") as fh:
        b64 = base64.b64encode(fh.read()).decode("ascii")
    via_b64 = er.scan_template(file_content_base64=b64, file_name="uploaded.xlsx")

    assert via_b64["fingerprint"] == via_path["fingerprint"]
    assert via_b64["summary"] == via_path["summary"]
    assert via_b64["tab_roles"] == via_path["tab_roles"]
    assert via_b64["formula_classification"] == via_path["formula_classification"]
    # base64 mode named the template from file_name; and it's a real scan
    assert via_b64["template_name"] == "uploaded"
    assert via_path["summary"]["tier_1b_blocks"] >= 1


def test_scan_requires_exactly_one_input():
    with pytest.raises(er.ExcelReportError):
        er.scan_template()                                        # neither
    with pytest.raises(er.ExcelReportError):
        er.scan_template(file_path="x.xlsx", file_content_base64="AA==")  # both


# --- inbox: bare filename resolution + listing -------------------------------

def test_inbox_lists_only_workbooks_newest_first(monkeypatch, tmp_path):
    monkeypatch.setattr(er, "EXCEL_INBOX_DIR", tmp_path)
    (tmp_path / "notes.txt").write_text("ignore me")
    (tmp_path / "old.xlsx").write_bytes(b"x")
    (tmp_path / "new.xlsx").write_bytes(b"x")
    os.utime(tmp_path / "old.xlsx", (1_000_000, 1_000_000))   # make "old" older

    res = er.list_inbox()
    names = [f["file_name"] for f in res["files"]]
    assert names == ["new.xlsx", "old.xlsx"]      # newest first, .txt filtered out
    assert res["count"] == 2
    assert res["inbox_path"] == str(tmp_path)


def test_bare_filename_resolves_against_inbox(monkeypatch, tmp_path):
    monkeypatch.setattr(er, "EXCEL_INBOX_DIR", tmp_path)
    wb = tmp_path / "Template.xlsx"
    wb.write_bytes(b"x")
    assert er._resolve_input_path("Template.xlsx") == str(wb)      # bare name -> inbox
    assert er._resolve_input_path(str(wb)) == str(wb)              # full path unchanged
    assert er._resolve_input_path("nope.xlsx") == "nope.xlsx"      # unknown left alone


def test_missing_file_error_points_at_the_inbox(monkeypatch, tmp_path):
    monkeypatch.setattr(er, "EXCEL_INBOX_DIR", tmp_path)
    with pytest.raises(er.ExcelReportError) as exc:
        er.scan_template(file_path="/mnt/user-data/uploads/Sandbox.xlsx")
    msg = str(exc.value)
    assert "sandbox" in msg.lower() and "list_excel_inbox" in msg   # self-explaining
