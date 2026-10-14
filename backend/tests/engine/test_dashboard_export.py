"""Excel export for live dashboard runs."""

from __future__ import annotations

import io

import openpyxl

from app.engine import dashboard_export as export


def test_build_workbook_has_summary_and_data_with_kpis_and_chart_panel():
    panels = [
        {
            "title": "Headlines",
            "chart_type": "kpi",
            "display": "kpi",
            "columns": ["METRIC", "VALUE", "FORMAT"],
            "rows": [
                ["Total Outstanding", 3_900_000, "currency"],
                ["Overdue Balance", 1_000_000, "currency"],
            ],
            "row_count": 2,
        },
        {
            "title": "Status Mix",
            "chart_type": "donut",
            "display": "chart",
            "orientation": None,
            "columns": ["STATUS", "BALANCE"],
            "rows": [["Outstanding", 100], ["Overdue", 50]],
            "row_count": 2,
        },
        {
            "title": "Detail",
            "chart_type": None,
            "display": "table",
            "columns": ["CLIENT", "AMOUNT"],
            "rows": [["Acme", 12_000]],
            "row_count": 1,
        },
    ]
    raw = export.build_dashboard_workbook("AR Health", panels, report_id="abc")
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    assert wb.sheetnames[0] == "Summary"
    assert "Data" in wb.sheetnames
    summary = wb["Summary"]
    assert summary["A1"].value == "AR Health"
    # KPI values landed on Summary
    metrics = {summary.cell(r, 1).value for r in range(1, 20)}
    assert "Total Outstanding" in metrics
    assert "Overdue Balance" in metrics
    data = wb["Data"]
    assert data["A1"].value == "Headlines"
    assert export.export_filename("AR Health!", "abc").endswith(".xlsx")
    assert "AR_Health" in export.export_filename("AR Health!", "abc")


def test_empty_panels_still_produce_workbook():
    raw = export.build_dashboard_workbook("Empty", [], report_id="x")
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    assert wb["Summary"]["A1"].value == "Empty"


def test_build_pdf_contains_title_and_kpis():
    panels = [
        {
            "title": "Headlines",
            "chart_type": "kpi",
            "display": "kpi",
            "columns": ["METRIC", "VALUE", "FORMAT"],
            "rows": [["Total Outstanding", 1000, "currency"]],
            "row_count": 1,
        },
        {
            "title": "Mix",
            "chart_type": "donut",
            "display": "chart",
            "columns": ["STATUS", "N"],
            "rows": [["Paid", 1], ["Overdue", 2]],
            "row_count": 2,
        },
    ]
    raw = export.build_dashboard_pdf(
        "AR Health",
        panels,
        report_id="r1",
        narrative="Outstanding is elevated versus last month.",
    )
    assert raw[:4] == b"%PDF"
    assert export.export_filename("AR", "r1", "pdf").endswith(".pdf")
