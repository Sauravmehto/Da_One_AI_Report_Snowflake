"""Unit tests for excel_reports' pure logic — query rendering, formula
re-extension, and workbook filling. No Snowflake connection required."""

import openpyxl
import pytest

from app.engine import excel_reports as er
from app.engine.sql_safety import UnsafeSQLError, validate_sql


# --- render_sql --------------------------------------------------------------

def test_render_sql_substitutes_and_escapes():
    sql = er.render_sql(
        "SELECT * FROM t WHERE yr = {{year}} AND region = {{region}}",
        [{"name": "year", "type": "number", "default": 2025},
         {"name": "region", "type": "string"}],
        {"region": "O'Brien"},
    )
    assert sql == "SELECT * FROM t WHERE yr = 2025 AND region = 'O''Brien'"


def test_render_sql_uses_default_when_value_missing():
    sql = er.render_sql("SELECT {{y}}", [{"name": "y", "type": "number", "default": 7}], {})
    assert sql == "SELECT 7"


def test_render_sql_missing_value_no_default_raises():
    with pytest.raises(er.ExcelReportError):
        er.render_sql("SELECT {{y}}", [{"name": "y", "type": "number"}], {})


def test_render_sql_rejects_non_numeric_number():
    with pytest.raises(er.ExcelReportError):
        er.render_sql("SELECT {{y}}", [{"name": "y", "type": "number"}], {"y": "1; DROP TABLE t"})


def test_render_sql_unfilled_placeholder_raises():
    with pytest.raises(er.ExcelReportError):
        er.render_sql("SELECT {{a}} {{b}}", [{"name": "a", "type": "number", "default": 1}], {})


def test_rendered_sql_still_gated_by_validate_sql():
    # A well-formed SELECT passes the same gate run_report relies on...
    ok = er.render_sql("SELECT * FROM t WHERE n = {{n}}", [{"name": "n", "type": "number"}], {"n": 5})
    assert validate_sql(ok)  # does not raise
    # ...and a non-SELECT is rejected, so the run path can never mutate data.
    with pytest.raises(UnsafeSQLError):
        validate_sql("UPDATE t SET x = 1")


# --- regenerate_formula ------------------------------------------------------

def test_regenerate_formula_reextends_range_end():
    assert er.regenerate_formula("=SUM(Data!B2:B100)", 41) == "=SUM(Data!B2:B41)"
    assert er.regenerate_formula("=SUMIFS($A$2:$A$99,$B$2:$B$99,1)", 50) == \
        "=SUMIFS($A$2:$A$50,$B$2:$B$50,1)"


def test_regenerate_formula_leaves_whole_column_range():
    assert er.regenerate_formula("=SUM(Data!B:B)", 41) == "=SUM(Data!B:B)"


# --- fill_workbook -----------------------------------------------------------

def test_fill_workbook_writes_rows_clears_stale_and_regenerates():
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    data["A1"] = "Region"
    data["B1"] = "Amount"
    # stale rows that must be cleared
    for i in range(10):
        data.cell(row=2 + i, column=1).value = "OLD"
        data.cell(row=2 + i, column=2).value = -1

    summary = wb.create_sheet("Summary")
    summary["B2"] = "=SUM(Data!B2:B100)"

    result = er.fill_workbook(
        wb, "Data",
        ["REGION", "AMOUNT"],
        [["East", 5], ["West", 9], ["North", 3]],
        [{"sheet": "Summary", "cell": "B2", "formula": "=SUM(Data!B2:B100)", "ranges": ["B2:B100"]}],
        header_row=1,
    )

    assert result["new_last_row"] == 4  # header row 1 + 3 data rows
    assert data["A2"].value == "East" and data["B4"].value == 3
    assert data["A5"].value is None      # stale row cleared
    assert summary["B2"].value == "=SUM(Data!B2:B4)"
    assert result["regenerated"][0]["to"] == "=SUM(Data!B2:B4)"


def test_fill_workbook_unknown_data_sheet_raises():
    wb = openpyxl.Workbook()
    with pytest.raises(er.ExcelReportError):
        er.fill_workbook(wb, "Nope", ["A"], [[1]], [], header_row=1)


def test_build_regeneration_plan_pulls_tier_1_only():
    classification = {
        "tier_1": [{"sheet": "S", "cell": "B2", "formula": "=SUM(D!B2:B9)", "ranges": ["B2:B9"]}],
        "tier_2": [{"sheet": "S", "cell": "B3", "formula": "=D!B5"}],
        "tier_3": [{"sheet": "S", "cell": "B4", "formula": "=B2*2"}],
    }
    plan = er.build_regeneration_plan(classification)
    assert len(plan) == 1 and plan[0]["cell"] == "B2"
