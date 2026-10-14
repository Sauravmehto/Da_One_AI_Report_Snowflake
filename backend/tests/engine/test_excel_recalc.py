"""The generated report must recalculate when opened.

The point of shipping a formula workbook rather than a values dump is that the
recipient opens it and sees live, correct numbers. openpyxl writes formulas
without evaluating them, so every formula cell leaves with no cached result — a
report whose totals render blank is worse than one that fails loudly.

The case that matters is a LOADED template, not a new workbook. A fresh
`Workbook()` already defaults to fullCalcOnLoad=True, so testing against one
proves nothing. `run_report` opens a workbook the user uploaded, and a loaded
workbook inherits whatever `<calcPr>` that file happens to carry — including
fullCalcOnLoad="0", which Excel does write.
"""

import io
import re
import zipfile

import openpyxl
import pytest

from app.engine.excel_reports import force_recalculation_on_open


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _calc_pr(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("xl/workbook.xml").decode()
    match = re.search(r"<calcPr[^>]*>", xml)
    return match.group(0) if match else ""


@pytest.fixture
def template_without_recalc() -> bytes:
    """A template as Excel may well have saved it: recalc-on-load switched off."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=SUM(B1:B9)"
    wb.calculation.fullCalcOnLoad = False
    return _save(wb)


def test_a_template_that_disables_recalc_would_poison_the_output(template_without_recalc):
    """Without the call, the flag propagates from the uploaded template into
    every report generated from it — and the totals may never compute."""
    loaded = openpyxl.load_workbook(io.BytesIO(template_without_recalc), data_only=False)

    assert loaded.calculation.fullCalcOnLoad is False
    assert 'fullCalcOnLoad="0"' in _calc_pr(_save(loaded))


def test_force_recalculation_overrides_the_template_setting(template_without_recalc):
    loaded = openpyxl.load_workbook(io.BytesIO(template_without_recalc), data_only=False)

    force_recalculation_on_open(loaded)

    assert 'fullCalcOnLoad="1"' in _calc_pr(_save(loaded))


def test_formulas_are_shipped_not_evaluated():
    """The documented consequence: formula cells carry no cached value, so
    preview-only tools (Outlook, SharePoint, Quick Look) show them blank until
    the file is opened in a spreadsheet application."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 10, 20, "=SUM(A1:A2)"
    force_recalculation_on_open(wb)
    data = _save(wb)

    with_formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    cached_only = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

    assert with_formulas.active["A3"].value == "=SUM(A1:A2)"
    assert cached_only.active["A3"].value is None
