"""Trim a filled US_RealEstate report download into an uploadable template."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

SRC = Path(__file__).resolve().parents[1] / "US_RealEstate_NATL_Report_Template.xlsx"
OUT = Path(__file__).resolve().parents[1] / "US_RealEstate_NATL_Template.xlsx"


def clear_rows_from(ws, start_row: int) -> None:
    last = ws.max_row or start_row
    if last >= start_row:
        ws.delete_rows(start_row, last - start_row + 1)


def rewrite_raw_ranges(formula: str, sheet: str, new_end: int = 3) -> str:
    """Point Raw_* ranges that grew to huge ends back to the sample end row."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    pat = re.compile(
        rf"({re.escape(sheet)}![A-Z]+\$?)2:([A-Z]+\$?)\d+",
        re.IGNORECASE,
    )
    return pat.sub(rf"\g<1>2:\g<2>{new_end}", formula)


def fix_formulas(ws, through_row: int) -> None:
    for r in range(1, through_row + 1):
        for c in range(1, 10):
            cell = ws.cell(r, c)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = rewrite_raw_ranges(cell.value, "Raw_Sales", 3)
                cell.value = rewrite_raw_ranges(cell.value, "Raw_Liens", 3)


def main() -> None:
    wb = openpyxl.load_workbook(SRC)

    for name in ("Raw_Sales", "Raw_Liens"):
        ws = wb[name]
        clear_rows_from(ws, 4)
        print(f"{name}: max_row={ws.max_row}")

    ws = wb["1_Sales_Summary"]
    fix_formulas(ws, 19)
    clear_rows_from(ws, 20)
    # Original template had a total under the two sample detail rows.
    ws["E20"] = "=SUM(E18:E19)"
    print(f"1_Sales_Summary: max_row={ws.max_row}")

    ws = wb["2_Lien_Activity"]
    fix_formulas(ws, 15)
    clear_rows_from(ws, 16)
    print(f"2_Lien_Activity: max_row={ws.max_row}")

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    wb.close()

    wb2 = openpyxl.load_workbook(OUT, read_only=True)
    for name in wb2.sheetnames:
        ws = wb2[name]
        print(f"OK {name}: rows={ws.max_row} cols={ws.max_column}")
    wb2.close()
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
